"""Descargas a Excel de los reportes de ventas."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = "Arial"
MONEDA = '$#,##0.00;-$#,##0.00;"-"'
CANTIDAD = '#,##0.###;-#,##0.###;"-"'


def libro_ventas_por_producto(context):
    """El reporte Ventas por producto (ver VentasPorProductoView) tal como se
    ve en pantalla: un renglón por producto o grupo y el total. Utilidad,
    margen y totales van como fórmulas, para que cuadren si se edita algo."""
    resultado = context["resultado"]
    por_producto = context["agrupar"] == "producto"

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas por producto"
    hoja["A1"] = "Ventas por producto"
    hoja["A1"].font = Font(name=FUENTE, size=14, bold=True)
    almacen = context["almacen"].nombre if context["almacen"] else "Todos"
    tipo = dict(context["tipos_venta"]).get(context["tipo"], "Ventas y remisiones")
    hoja["A2"] = (
        f"Del {context['fecha_desde']:%d/%m/%Y} al {context['fecha_hasta']:%d/%m/%Y} · Almacén: {almacen} · "
        f"{tipo} · Agrupado por: {context['agrupar_etiqueta']}"
    )
    hoja["A2"].font = Font(name=FUENTE, size=9, italic=True)

    columnas = [(context["agrupar_etiqueta"], 42, None)]
    if por_producto:
        columnas.append(("SKU", 14, None))
    columnas += [("Registros", 11, "#,##0"), ("Cantidad", 13, CANTIDAD), ("Cantidad auxiliar", 15, CANTIDAD)]
    if por_producto:
        columnas.append(("Unidad auxiliar", 12, None))
    columnas += [
        ("Costo", 15, MONEDA), ("Importe sin impuesto", 17, MONEDA), ("Utilidad", 15, MONEDA),
        ("Importe con impuesto", 17, MONEDA), ("Margen", 10, "0.0%"),
    ]
    letra = {titulo: get_column_letter(i) for i, (titulo, _, _) in enumerate(columnas, start=1)}
    ENCABEZADO = 4
    for i, (titulo, ancho, _) in enumerate(columnas, start=1):
        celda = hoja.cell(ENCABEZADO, i, titulo)
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="334155")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = ancho

    costo, importe, utilidad = letra["Costo"], letra["Importe sin impuesto"], letra["Utilidad"]

    def valores(fila, r):
        datos = {
            context["agrupar_etiqueta"]: fila["nombre"], "SKU": fila.get("secundario", ""),
            "Registros": fila["registros"], "Cantidad": float(fila["cantidad"]),
            "Cantidad auxiliar": float(fila["cantidad_auxiliar"]), "Unidad auxiliar": fila.get("unidad_auxiliar", ""),
            "Costo": float(fila["costo"]), "Importe sin impuesto": float(fila["importe"]),
            "Utilidad": f"={importe}{r}-{costo}{r}", "Importe con impuesto": float(fila["total"]),
            "Margen": f'=IF({importe}{r}=0,"",{utilidad}{r}/{importe}{r})',
        }
        return [datos[titulo] for titulo, _, _ in columnas]

    primera = ENCABEZADO + 1
    for r, fila in enumerate(resultado["filas"], start=primera):
        for i, valor in enumerate(valores(fila, r), start=1):
            celda = hoja.cell(r, i, valor)
            celda.font = Font(name=FUENTE, size=10)
            if columnas[i - 1][2]:
                celda.number_format = columnas[i - 1][2]

    ultima = primera + len(resultado["filas"]) - 1
    total = ultima + 1
    hoja.cell(total, 1, "Total")
    for i, (titulo, _, formato) in enumerate(columnas, start=1):
        col = get_column_letter(i)
        if titulo == "Margen":
            hoja.cell(total, i, f'=IF({importe}{total}=0,"",{utilidad}{total}/{importe}{total})')
        elif formato and titulo != "Unidad auxiliar" and i > 1:
            hoja.cell(total, i, f"=SUM({col}{primera}:{col}{ultima})" if ultima >= primera else 0)
        celda = hoja.cell(total, i)
        celda.font = Font(name=FUENTE, size=10, bold=True)
        if formato:
            celda.number_format = formato
    hoja.freeze_panes = hoja.cell(primera, 2)
    if ultima >= primera:
        hoja.auto_filter.ref = f"A{ENCABEZADO}:{get_column_letter(len(columnas))}{ultima}"
    return libro
