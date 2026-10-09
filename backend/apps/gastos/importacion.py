"""Importación de los vales históricos del Excel de pólizas de gastos.

Se hace en dos pasos para que una persona revise antes de que nada entre a
la base (los gastos son datos reales):

1. `generar_revision_gastos` lee el Excel, propone por cada vale qué hacer
   (IMPORTAR, OMITIR o REVISAR, con el motivo) y con qué fecha, centro de
   costo y concepto, y escribe un Excel de revisión.
2. `importar_gastos_revisados` lee ese Excel ya revisado, valida todas las
   filas marcadas IMPORTAR y, solo si ninguna tiene error, las crea.

Reglas acordadas con el negocio para la propuesta:
- Un número de vale repetido con el mismo importe, descripción y fecha se
  importa una sola vez; si los importes difieren (partes de un vale o
  reclasificaciones entre pólizas) todas sus filas quedan para revisar.
- "Bodega-Lerdo" (casi siempre la policía de ambas) se carga completo a
  Bodega Sur. "Sucursales"/"Varias sucursales" quedan para revisar.
- La mercancía para venta y la hoja de compra de facturas se proponen
  omitir. Un vale sin concepto identificable queda para revisar.
- Los vales de hojas o secciones fiscales se marcan facturados con el folio
  pendiente (el Excel no trae el UUID del proveedor).
- Si la fecha es posterior al Excel y la sección dice otro año, se corrige
  el año con el de la sección (errores de captura como 28/11/2026 en una
  sección de noviembre 2025).
"""
import datetime
import difflib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

IMPORTAR, OMITIR, REVISAR = "IMPORTAR", "OMITIR", "REVISAR"
ACCIONES = (IMPORTAR, OMITIR, REVISAR)
FOLIO_PENDIENTE = "Folio pendiente (importado)"

MESES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8,
    "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}
RE_FECHA = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{2,4})")
RE_ANIO = re.compile(r"20\d\d")


def abrir_libro(ruta, read_only=False):
    """Carga el Excel desde memoria: openpyxl deja el archivo abierto (y
    bloqueado en Windows) hasta que se libera el objeto, lo que impide
    volver a guardarlo en Excel mientras corre el comando."""
    with open(ruta, "rb") as archivo:
        contenido = io.BytesIO(archivo.read())
    return openpyxl.load_workbook(contenido, data_only=True, read_only=read_only)


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().upper()
    return " ".join(texto.replace(".", " ").split())


# --------------------------------------------------------------------------
# Lectura del Excel de pólizas
# --------------------------------------------------------------------------

@dataclass
class ValeExcel:
    hoja: str
    fila: int
    seccion: str
    vale: str
    descripcion: str
    centro_original: str
    importe: Decimal
    condicion: str
    turno_anterior: str
    factura_caja: str
    folio_factura: str
    fecha: datetime.date | None
    mes_seccion: int | None
    anio_seccion: int | None

    @property
    def clave(self):
        return f"{self.hoja}!{self.fila}"


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _fecha_de_texto(texto):
    m = RE_FECHA.search(texto or "")
    if not m:
        return None
    dia, mes, anio = (int(x) for x in m.groups())
    if anio < 100:
        anio += 2000
    try:
        return datetime.date(anio, mes, dia)
    except ValueError:
        return None


def _mes_y_anio(seccion):
    texto = normalizar(seccion)
    mes = next((num for nombre, num in MESES.items() if nombre in texto), None)
    anio = RE_ANIO.search(texto)
    return mes, int(anio.group()) if anio else None


def _es_seccion(celdas):
    """Encabezado de una póliza, p. ej. "GASTOS DE SUCURSALES SIN FACTURA MES
    DE MARZO 2026": solo texto, y menciona gastos, facturas o un mes. Las
    notas sueltas ("SE PASO") no cambian la sección vigente."""
    textos = [c for c in celdas if isinstance(c, str) and c.strip()]
    if not textos or any(isinstance(c, (int, float)) and not isinstance(c, bool) for c in celdas):
        return None
    texto = normalizar(textos[0])
    if "GASTO" in texto or "FACTURA" in texto or any(mes in texto for mes in MESES):
        return textos[0].strip()
    return None


def leer_polizas(ruta):
    """Todos los vales del Excel, de todas las hojas. Una fila es vale si
    tiene descripción e importe bajo un encabezado con CONCEPTO e IMPORTE;
    los subtotales (sin descripción) y las tablas de facturas (encabezado
    sin CONCEPTO) se ignoran."""
    libro = abrir_libro(ruta, read_only=True)
    try:
        return _leer_hojas(libro)
    finally:
        libro.close()


def _leer_hojas(libro):
    vales = []
    for hoja in libro.worksheets:
        columnas = None
        seccion = ""
        for numero, fila in enumerate(hoja.iter_rows(values_only=True, max_col=14), start=1):
            celdas = list(fila)
            if not any(c is not None and str(c).strip() for c in celdas):
                continue
            nombres = [normalizar(c) if isinstance(c, str) else "" for c in celdas]
            if "IMPORTE" in nombres:
                columnas = nombres if "CONCEPTO" in nombres else None
                continue
            nueva_seccion = _es_seccion(celdas)
            if nueva_seccion:
                seccion = nueva_seccion
                continue
            if columnas is None:
                continue

            def valor(*nombres_columna):
                for nombre in nombres_columna:
                    if nombre in columnas:
                        return celdas[columnas.index(nombre)]
                return None

            descripcion = valor("CONCEPTO")
            importe = valor("IMPORTE")
            if not isinstance(descripcion, str) or len(descripcion.strip()) < 4:
                continue
            if not isinstance(importe, (int, float)) or isinstance(importe, bool):
                continue
            condicion = valor("CONDICION")
            if isinstance(condicion, datetime.datetime):
                fecha = condicion.date()
                condicion_texto = fecha.strftime("%d/%m/%Y")
            else:
                condicion_texto = _texto(condicion)
                fecha = _fecha_de_texto(condicion_texto)
            mes, anio = _mes_y_anio(seccion)
            vale = valor("VALE", "FOLIO")
            vales.append(ValeExcel(
                hoja=hoja.title.strip(),
                fila=numero,
                seccion=seccion,
                vale="" if isinstance(vale, datetime.datetime) else _texto(vale),
                descripcion=" ".join(descripcion.split()),
                centro_original=_texto(valor("CENTRO DE COSTO", "SUCURSAL")),
                importe=Decimal(str(importe)).quantize(Decimal("0.01")),
                condicion=condicion_texto,
                turno_anterior=_texto(valor("TURNO")),
                factura_caja=_texto(valor("FACTURA DONDE SALE EL GASTO")),
                folio_factura=_texto(valor("FOLIO FACTURA")),
                fecha=fecha,
                mes_seccion=mes,
                anio_seccion=anio,
            ))
    return vales


# --------------------------------------------------------------------------
# Reglas de la propuesta
# --------------------------------------------------------------------------

ALIAS_CENTRO = {
    "BODEGA SUR": "BSU", "B SUR": "BSU", "BODEGA": "BSU",
    "LERDO": "LER", "ALDAMA": "ALD", "TECOLUTILLA": "TEC", "IQUINUAPA": "IQU", "MERCADO": "MER",
    "HUIMANGUILLO": "HUI",
    "ADMINISTRACION": "ADM", "ADMON": "ADM", "ADMINISTRATIVO": "ADM",
    "GASTO PERSONAL": "PER", "GASTOS PERSONALES": "PER", "GASTOS PERSONAL": "PER", "GAS PERSONAL": "PER",
    "PERSONAL": "PER",
    "CASA PLAYA": "CPL", "CASA DE LA PLAYA": "CPL", "PLAYA": "CPL",
    "SALON DE JUNTAS": "SJU",
    "MOBILIARIO": "MOB", "TRANSPORTES": "TRA", "TRANSPORTE": "TRA",
    "ARRENDAMIENTO": "ARR", "GASTO DE ARRENDAMIENTO": "ARR", "RENTAS": "ARR",
    "SEMBRANDO VIDA": "PSV", "PROYECTO VIGALTA": "PVI", "VIGALTA": "PVI",
    "CITRICO": "PCI", "CITRICOS": "PCI", "PROYECTO JARDIN": "PJA",
    "PROYECTO BECERRO": "PBE", "PROYECTO BECERROS": "PBE",
    "PROYECTO BOVINO": "PBO", "BOVINO": "PBO", "PROY BOVINO": "PBO", "PROYECTO BOV": "PBO",
    "PROYECTO CACAO": "PCA", "RANCHO": "RAN", "GASTO RANCHO": "RAN", "GASTOS RANCHO": "RAN",
}
SIN_CENTRO_PRECISO = {"SUCURSALES", "VARIAS SUCURSALES", "TODAS SUCURSALES", "VARIAS", "SUC VARIAS", "VAR SUC"}


def proponer_centro(centro_original):
    """(código, nota) o (None, motivo para revisar)."""
    texto = normalizar(centro_original)
    texto = re.sub(r"^SUC\b[ -]*", "", texto).replace(" - ", "-").replace("- ", "-").replace(" -", "-")
    if not texto:
        return None, "El vale no trae centro de costo."
    if "BODEGA" in texto and "LERDO" in texto and texto.count("-") <= 1:
        return "BSU", "Compartido Bodega Sur-Lerdo: se carga completo a Bodega Sur."
    if texto in ALIAS_CENTRO:
        return ALIAS_CENTRO[texto], ""
    if texto in SIN_CENTRO_PRECISO:
        return None, f"Centro «{centro_original}» sin sucursal precisa: asignar uno."
    parecido = difflib.get_close_matches(texto, [*ALIAS_CENTRO, *SIN_CENTRO_PRECISO], n=1, cutoff=0.82)
    if parecido and parecido[0] in ALIAS_CENTRO:
        return ALIAS_CENTRO[parecido[0]], f"Centro «{centro_original}» leído como {parecido[0]}."
    if parecido:
        return None, f"Centro «{centro_original}» sin sucursal precisa: asignar uno."
    return None, f"Centro «{centro_original}» no reconocido: asignar uno."


HOJAS_PROYECTO = {"SEMBRANDO VIDA", "PROYECTO VIGALTA", "CITRICOS", "JARDIN", "PROYECTO BECERROS",
                  "PROYECTO BOVINO-GANADO", "CDIS", "PROYECTO CACAO"}
RE_JORNAL = re.compile(
    r"\bTRAB(AJO|AJOS|J)?\b|\bPOR (TRA|TBA)\b|LIMPIEZA DEL|LIMPIEZA DE (TERRENO|RANCHO|POTRERO)|CHAPE|JORNAL|"
    r"DIAS DE TRABAJO"
)

# (expresión sobre la descripción normalizada, concepto). Gana la primera que
# coincide, así que las más específicas van primero.
REGLAS_CONCEPTO = [
    (r"POLICIA|POLCIA|POLICIAS|VIGILAN|RONDIN", "Servicio de Vigilancia"),
    (r"CONSUMO INTERNO|CONSUMO.*\b(GATOS?|PERROS?|MASCOTAS?|AVES|VES)\b|CONSUMO (DE )?SUC\b|DESPARA",
     "Consumo Interno de Mercancía (animales)"),
    (r"MANIOBRA|DESCARGA|CARGADOR", "Maniobras y descarga de mercancía"),
    (r"CASETA|PEAJE", "Peajes de caseta"),
    (r"VIATICO", "Viáticos y Gastos de Viaje"),
    (r"FINIQUITO|INDEMNIZ", "Indemnización"),
    (r"AGUINALDO", "Aguinaldo"),
    (r"NOMINA|\bQNA\b|QUINCENA|PAGO DE SEM\b|SUELDO", "Sueldos y Salarios"),
    (r"COMISION.*(BANC|OXXO)|(BANC|OXXO).*COMISION", "Comisiones Bancarias"),
    (r"GASOLINA|DIESEL|COMBUSTIBLE|\bMAGNA\b|PREMIUM|ACEITE", "Combustibles y Lubricantes"),
    (r"FLETE|ACARREO|VIAJES? DE (TIERRA|RELLENO|ARENA|GRAVA|PIEDRA)|PAQUETERIA|\bENVIO\b", "Fletes y Acarreos"),
    (r"\bLUZ\b|\bCFE\b|ENERGIA ELECTRICA", "Servicio de Energía Eléctrica"),
    (r"\bAGUA\b|GARRAF|BONAFON", "Agua"),
    (r"INTE?RNET|TELMEX|IZZI|TOTALPLAY", "Servicio de Internet"),
    (r"TELCEL|RECARGA|CELULAR|AT&T|ATT\b", "Servicio de Telefonía Móvil"),
    (r"PAGO DE RENTA|RENTA DE|RENTA SUC|\bRENTA\b", "Arrendamientos a Personas Físicas"),
    (r"PREDIAL", "Impuesto Predial"),
    (r"REFRENDO|TENENCIA", "Refrendos"),
    (r"\bIMSS\b", "Cuota IMSS"),
    (r"INFONAVIT", "Aportaciones al Infonavit"),
    (r"FUMIGA|PLAGA|RATICIDA", "Servicios de limpieza especializada (fumigación)"),
    (r"LLANTA|REFACC|BALATA|AMORTIGUADOR", "Refacciones de equipo de transporte"),
    (r"AFINACION|MECANIC|ALINEACION|LAVADO GENERAL|ENGRASADO|PARCHADURA|HOJALATER", "Mantenimiento de equipo de transporte"),
    (r"(REPAR|MANT+[EO0]|SERVICIO).*(COSTURADORA|MONTACARGA|MAQUINA|DESBROZADORA|MOTOSIERRA|PODADORA|PLANTA|TRACTOR)",
     "Mantenimiento de maquinaria y equipo"),
    (r"MINISPLIT|AIRE ACONDICIONADO|\bCLIMA\b", "Mantenimiento de climatización"),
    (r"PLOMER|TUBERIA|\bFUGA\b|\bBANOS?\b|\bWC\b|TINACO|DRENAJE|MEZCLADORA", "Mantenimiento hidráulico y sanitario"),
    (r"ELECTRIC|\bFOCOS?\b|LAMPARA|BREAKER|APAGADOR|CONTACTO", "Mantenimiento eléctrico"),
    (r"PINTURA|IMPERMEAB|ALBANIL|CEMENTO|\bBLOCK|VARILLA|TECHO|\bPISO|PUERTA|CERRADURA|CHAPA\b|LLAVES|HERRERIA|"
     r"ALUMINI|SOLDA|CALIDRA|MORTERO|CONS?TRUCCION|CARPINTERIA|ALBERCA",
     "Mantenimiento y Conservación de edificios e instalaciones"),
    (r"TONER|CARTUCHO|\bTINTA", "Consumibles de Impresión"),
    (r"MOUSE|TECLADO|\bUSB\b|DISCO DURO|MEMORIA|\bPILAS\b", "Accesorios y componentes de equipo de cómputo"),
    (r"PAPELERIA|\bHOJAS\b|LIBRETA|PLUMAS?\b|COPIAS|FOLDER|ENGARGOL", "Papelería y Artículos de Oficina"),
    (r"PAPEL HIGIENICO|CLORO|ESCOBA|TRAPEADOR|JABON|DETERGENTE|PINOL|FABULOSO|ARTICULOS DE LIMPIEZA|"
     r"MATERIAL DE LIMPIEZA|BOLSAS? DE BASURA", "Artículos de Limpieza"),
    (r"BOLSA|EMPAQUE|PLAYO|EMPLAYE|COSTAL|HILO PARA LA COSTURADORA|CINTA CANELA", "Empaques y Envolturas"),
    (r"UNIFORME|PLAYERA|CAMISA", "Uniformes"),
    (r"SUERO|BOTIQUIN|ELECTROLIT|GASAS|PRIMEROS AUXILIOS", "Salud Ocupacional"),
    (r"REGALO|CORTESIA|ATENCION A CLIENTE|PARA LOS CLIENTE", "Atención a Clientes"),
    (r"PUBLICIDAD|\bSPOTS?\b|\bLONAS?\b|VOLANTE|PERIFONEO", "Propaganda y Publicidad"),
    (r"HONORARIO|ASESORIA|CONTADOR|ABOGADO|NOTARI", "Honorarios a Personas Físicas"),
    (r"CAPACITACION|\bCURSO\b", "Capacitación a Empleados"),
    (r"HERRAMIENTA|CABOS? DE|\bCOA\b|H?ACHA\b|MARTILLO|PINZAS?\b|DESARMADOR|MACHETE|\bPALAS?\b|CARRETILLA|SEGUETA|LIMA PARA", "Herramientas de Trabajo"),
    (r"\bGAS\b|TANQUE DE GAS|CILINDRO", "Gas"),
    (r"MEDIC|FARMACIA|CONSULTA|DOCTOR|HOSPITAL|LABORATORIO|DENTISTA", "Gastos Médicos"),
    (r"COLEGIATURA|INSCRIPCION|ESCUELA", "Colegiatura"),
    # Al final: mantenimiento sin más detalle (los específicos ya pasaron).
    (r"MANT+[EO0]|MANTENIMIENTO|MATTO|MANNTTO|REPARA?C|ARREGLO|COMPOSTURA",
     "Mantenimiento y Conservación de edificios e instalaciones"),
]
# Solo para centros de tipo Personal (gastos de los dueños).
REGLAS_CONCEPTO_PERSONAL = [
    (re.compile(r"TARJETA"), "Tarjeta Personal"),
    (re.compile(r"EFEC(TIVO)? PARA (GASTO|PAGO)|\bTANDA\b|\bRETIRO\b|ANTICIPO"), "Retiros y movimientos de la propietaria"),
]
REGLAS_CONCEPTO = [(re.compile(patron), concepto) for patron, concepto in REGLAS_CONCEPTO]

RE_MERCANCIA_VENTA = re.compile(r"PARA (LA )?VENTA|PARA VTA|\bVTA\b|TEMU")

# Unidades del catálogo de vehículos que se reconocen en la descripción.
REGLAS_VEHICULO = [
    (re.compile(p), nombre) for p, nombre in [
        (r"\bKW[ -]?2010|KENWORTH 2010", "Kenworth 2010"), (r"\bKW[ -]?(20)?17|KENWORTH 2017", "Kenworth 2017"),
        (r"ISUZU 2018", "Isuzu 2018"), (r"ISUZU 2023", "Isuzu 2023"),
        (r"NISSAN 2011", "Nissan 2011"), (r"NISSAN 2014", "Nissan 2014"), (r"NISSAN 2016", "Nissan 2016"),
        (r"NISSAN 2023", "Nissan 2023"), (r"MARCH 2019", "March 2019"), (r"MARCH 2023", "March 2023"),
        (r"MOTO ?CARRO", "Motocarro"), (r"MONTACARGA", "Montacargas"), (r"INTERNATIONAL|INTERNACIONAL", "International"),
        (r"TRACTOR", "Tractor"), (r"DESBROZADORA", "Desbrozadora"), (r"MOTOSIERRA", "Motosierra"),
        (r"PODADORA", "Podadora"),
    ]
]


def proponer_concepto(vale, centro=""):
    descripcion = normalizar(vale.descripcion)
    if normalizar(vale.hoja) in HOJAS_PROYECTO and RE_JORNAL.search(descripcion):
        return "Mano de obra eventual (jornales)"
    if centro == "PER":
        for patron, concepto in REGLAS_CONCEPTO_PERSONAL:
            if patron.search(descripcion):
                return concepto
    for patron, concepto in REGLAS_CONCEPTO:
        if patron.search(descripcion):
            return concepto
    return None


def proponer_vehiculo(vale):
    descripcion = normalizar(vale.descripcion)
    return next((nombre for patron, nombre in REGLAS_VEHICULO if patron.search(descripcion)), "")


def es_facturado(vale):
    hoja, seccion = normalizar(vale.hoja), normalizar(vale.seccion)
    if vale.folio_factura:
        return True
    if re.search(r"\bSF\b|S/F|SIN FACTURA|NO FACTURADO", f"{hoja} {seccion}"):
        return False
    return bool(re.search(r"FISCAL|FACTURADOS|CON FACTURA", f"{hoja} {seccion}"))


@dataclass
class Propuesta:
    vale: ValeExcel
    accion: str = IMPORTAR
    motivos: list = field(default_factory=list)
    fecha: datetime.date | None = None
    centro: str = ""
    concepto: str = ""
    vehiculo: str = ""
    facturado: bool = False
    referencia_factura: str = ""

    def marcar(self, accion, motivo):
        # OMITIR pesa más que REVISAR, y ambos más que IMPORTAR.
        orden = {IMPORTAR: 0, REVISAR: 1, OMITIR: 2}
        if orden[accion] > orden[self.accion]:
            self.accion = accion
        self.motivos.append(motivo)


def _fecha_propuesta(propuesta, fecha_corte):
    vale = propuesta.vale
    fecha = vale.fecha
    if fecha is None:
        if vale.mes_seccion and vale.anio_seccion:
            propuesta.marcar(REVISAR, "Sin fecha en el vale: se propone el día 1 del mes de la sección.")
            return datetime.date(vale.anio_seccion, vale.mes_seccion, 1)
        propuesta.marcar(REVISAR, "Sin fecha en el vale ni en la sección.")
        return None
    if fecha > fecha_corte:
        if vale.anio_seccion and vale.anio_seccion != fecha.year:
            try:
                corregida = fecha.replace(year=vale.anio_seccion)
            except ValueError:
                corregida = None
            if corregida and corregida <= fecha_corte:
                propuesta.motivos.append(
                    f"Año corregido de {fecha.year} a {vale.anio_seccion} según la sección (fecha posterior al Excel)."
                )
                return corregida
        propuesta.marcar(REVISAR, "Fecha posterior a la del Excel.")
    return fecha


def proponer(vales, anio, fecha_corte, conceptos_validos, centros_validos, vehiculos_validos, conceptos_personales,
             centros_personales):
    """Propuestas para los vales que caen en `anio`, ya con los repetidos
    resueltos. Los catálogos se reciben como conjuntos de nombres/códigos
    para no consultar la base aquí."""
    propuestas = []
    for vale in vales:
        propuesta = Propuesta(vale=vale)
        propuesta.fecha = _fecha_propuesta(propuesta, fecha_corte)
        fecha_para_anio = propuesta.fecha or (
            datetime.date(vale.anio_seccion, 1, 1) if vale.anio_seccion else None
        )
        if fecha_para_anio is None or fecha_para_anio.year != anio:
            continue

        hoja = normalizar(vale.hoja)
        descripcion = normalizar(vale.descripcion)
        if hoja == "COMPRA DE FACT":
            propuesta.marcar(OMITIR, "Compra de facturas: se define con la contadora.")
        if hoja == "VTA REFRESCO" or RE_MERCANCIA_VENTA.search(descripcion):
            propuesta.marcar(OMITIR, "Mercancía para venta: va por Compras/Inventario, no es gasto.")
        if "AMPARA" in normalizar(vale.seccion):
            propuesta.marcar(OMITIR, "Fila de una factura que ampara vales, no es un vale.")

        centro, nota = proponer_centro(vale.centro_original)
        if centro and centro not in centros_validos:
            centro, nota = None, f"El centro {centro} no existe en el sistema: asignar uno."
        if centro:
            propuesta.centro = centro
            if nota:
                propuesta.motivos.append(nota)
        else:
            propuesta.marcar(REVISAR, nota)

        concepto = proponer_concepto(vale, propuesta.centro)
        if concepto and concepto not in conceptos_validos:
            raise ValueError(f"La regla propone el concepto «{concepto}», que no existe en el catálogo.")
        if concepto is None:
            propuesta.marcar(REVISAR, "Sin concepto identificable: asignar uno.")
        else:
            propuesta.concepto = concepto
            if concepto in conceptos_personales and centro and centro not in centros_personales:
                propuesta.marcar(REVISAR, f"«{concepto}» es gasto personal y el centro no es de tipo Personal.")

        vehiculo = proponer_vehiculo(vale)
        propuesta.vehiculo = vehiculo if vehiculo in vehiculos_validos else ""
        propuesta.facturado = es_facturado(vale)
        if propuesta.facturado:
            propuesta.referencia_factura = vale.folio_factura or FOLIO_PENDIENTE
        propuestas.append(propuesta)

    _resolver_repetidos(propuestas)
    return propuestas


def _resolver_repetidos(propuestas):
    por_vale = {}
    for propuesta in propuestas:
        if propuesta.vale.vale:
            por_vale.setdefault(propuesta.vale.vale, []).append(propuesta)
    for numero, grupo in por_vale.items():
        if len(grupo) < 2:
            continue
        firmas = {(p.vale.importe, normalizar(p.vale.descripcion), p.fecha) for p in grupo}
        if len(firmas) == 1:
            original = grupo[0]
            for repetida in grupo[1:]:
                repetida.marcar(OMITIR, f"Duplicado exacto de {original.vale.clave}.")
            continue
        importes = ", ".join(f"${p.vale.importe} ({p.vale.clave})" for p in grupo)
        for propuesta in grupo:
            propuesta.marcar(
                REVISAR, f"El vale {numero} aparece {len(grupo)} veces con datos distintos: {importes}. "
                "Decidir cuáles importar.",
            )


# --------------------------------------------------------------------------
# Archivo de revisión
# --------------------------------------------------------------------------

COLUMNAS = [
    ("clave", "CLAVE (no editar)", 26),
    ("accion", "ACCIÓN", 12),
    ("motivo", "MOTIVO", 60),
    ("fecha", "FECHA", 12),
    ("vale", "VALE", 10),
    ("descripcion", "DESCRIPCIÓN", 50),
    ("importe", "IMPORTE", 12),
    ("centro", "CENTRO DE COSTO", 26),
    ("concepto", "CONCEPTO", 40),
    ("vehiculo", "VEHÍCULO", 16),
    ("facturado", "FACTURADO", 11),
    ("referencia_factura", "REFERENCIA FACTURA", 26),
    ("hoja", "HOJA", 22),
    ("seccion", "SECCIÓN", 40),
    ("centro_original", "CENTRO EN EXCEL", 20),
    ("condicion", "CONDICIÓN EN EXCEL", 22),
    ("turno_anterior", "TURNO ANTERIOR", 12),
    ("factura_caja", "FACTURA DE CAJA", 14),
]
COLORES = {IMPORTAR: "E8F5E9", REVISAR: "FFF8E1", OMITIR: "ECEFF1"}
INSTRUCCIONES = [
    "Archivo de revisión de vales de gastos.",
    "",
    "1. Revisa la hoja Vales. Cada fila trae lo que el sistema propone y, en MOTIVO, por qué.",
    "2. ACCIÓN: IMPORTAR la mete al sistema; OMITIR la descarta; REVISAR no se importa todavía.",
    "   Al terminar no debe quedar ninguna fila en REVISAR que quieras importar: cámbiala a IMPORTAR u OMITIR.",
    "3. Puedes corregir FECHA, DESCRIPCIÓN, IMPORTE, CENTRO DE COSTO, CONCEPTO, VEHÍCULO, FACTURADO y",
    "   REFERENCIA FACTURA. Centro, concepto y vehículo se eligen de la lista (hoja Listas).",
    "4. No cambies la columna CLAVE: identifica la fila del Excel original y evita importar dos veces.",
    "5. Si un vale se debe repartir entre varios centros, divídelo en varias filas con la misma CLAVE",
    "   seguida de /1, /2... (p. ej. GASTOS SF SUCURSALES!16080/1) y el importe de cada parte.",
    "",
    "Después se corre: python manage.py importar_gastos_revisados <este archivo> (solo valida)",
    "y, si no hay errores: python manage.py importar_gastos_revisados <este archivo> --aplicar",
]


def escribir_revision(ruta, propuestas, centros, conceptos, vehiculos):
    """`centros`: [(código, nombre)], `conceptos`: [(grupo, nombre)],
    `vehiculos`: [nombre]."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Vales"
    hoja.append([titulo for _, titulo, _ in COLUMNAS])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for indice, (_, _, ancho) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[openpyxl.utils.get_column_letter(indice)].width = ancho
    etiqueta_centro = dict(centros)
    for propuesta in sorted(propuestas, key=lambda p: (p.fecha or datetime.date.max, p.vale.hoja, p.vale.fila)):
        vale = propuesta.vale
        hoja.append([
            vale.clave,
            propuesta.accion,
            " ".join(propuesta.motivos),
            propuesta.fecha,
            vale.vale,
            vale.descripcion,
            float(vale.importe),
            f"{propuesta.centro} - {etiqueta_centro[propuesta.centro]}" if propuesta.centro else "",
            propuesta.concepto,
            propuesta.vehiculo,
            "SI" if propuesta.facturado else "NO",
            propuesta.referencia_factura,
            vale.hoja,
            vale.seccion,
            vale.centro_original,
            vale.condicion,
            vale.turno_anterior,
            vale.factura_caja,
        ])
        relleno = PatternFill("solid", fgColor=COLORES[propuesta.accion])
        for celda in hoja[hoja.max_row]:
            celda.fill = relleno
        hoja.cell(row=hoja.max_row, column=4).number_format = "DD/MM/YYYY"
        hoja.cell(row=hoja.max_row, column=7).number_format = "#,##0.00"
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions

    listas = libro.create_sheet("Listas")
    listas.append(["CENTRO DE COSTO", "CONCEPTO", "GRUPO DEL CONCEPTO", "VEHÍCULO"])
    filas = max(len(centros), len(conceptos), len(vehiculos))
    for i in range(filas):
        listas.append([
            f"{centros[i][0]} - {centros[i][1]}" if i < len(centros) else None,
            conceptos[i][1] if i < len(conceptos) else None,
            conceptos[i][0] if i < len(conceptos) else None,
            vehiculos[i] if i < len(vehiculos) else None,
        ])
    for letra, ancho in (("A", 34), ("B", 50), ("C", 32), ("D", 18)):
        listas.column_dimensions[letra].width = ancho

    ultima = max(hoja.max_row, 2)
    validaciones = [
        ("B", f'"{",".join(ACCIONES)}"'),
        ("H", f"Listas!$A$2:$A${len(centros) + 1}"),
        ("I", f"Listas!$B$2:$B${len(conceptos) + 1}"),
        ("J", f"Listas!$D$2:$D${len(vehiculos) + 1}"),
        ("K", '"SI,NO"'),
    ]
    for columna, formula in validaciones:
        validacion = DataValidation(type="list", formula1=formula, allow_blank=True)
        validacion.add(f"{columna}2:{columna}{ultima}")
        hoja.add_data_validation(validacion)

    instrucciones = libro.create_sheet("Instrucciones", 0)
    for linea in INSTRUCCIONES:
        instrucciones.append([linea])
    instrucciones.column_dimensions["A"].width = 110
    libro.active = 1
    libro.save(ruta)


def leer_revision(ruta):
    """Filas del archivo de revisión como diccionarios (claves de COLUMNAS)."""
    libro = abrir_libro(ruta)
    try:
        return _leer_vales_revisados(libro)
    finally:
        libro.close()


def _leer_vales_revisados(libro):
    if "Vales" not in libro.sheetnames:
        raise ValueError("El archivo no tiene la hoja «Vales»; ¿es un archivo de revisión de gastos?")
    hoja = libro["Vales"]
    titulos = [normalizar(c.value) for c in hoja[1]]
    indices = {}
    for clave, titulo, _ in COLUMNAS:
        normal = normalizar(titulo)
        if normal not in titulos:
            raise ValueError(f"Falta la columna «{titulo}» en la hoja Vales.")
        indices[clave] = titulos.index(normal)
    filas = []
    for numero, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in fila):
            continue
        datos = {clave: fila[i] for clave, i in indices.items()}
        datos["fila_revision"] = numero
        filas.append(datos)
    return filas


def a_fecha(valor):
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def a_importe(valor):
    try:
        importe = Decimal(str(valor).replace(",", "").replace("$", "").strip()).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None
    return importe if importe > 0 else None


def codigo_de_centro(valor):
    """"BSU - Bodega Sur" -> "BSU"; acepta también solo el código."""
    return str(valor or "").split(" - ", 1)[0].strip().upper()


# --------------------------------------------------------------------------
# Importación del archivo revisado
# --------------------------------------------------------------------------

@dataclass
class ResultadoImportacion:
    gastos: list = field(default_factory=list)
    errores: list = field(default_factory=list)
    ya_importadas: int = 0
    por_accion: dict = field(default_factory=dict)
    creados: int = 0


def _recortar(texto, largo):
    texto = _texto(texto)
    return texto[:largo]


def preparar_gastos(filas):
    """Valida todas las filas marcadas IMPORTAR y arma sus gastos sin
    guardarlos. Las que ya se importaron antes (misma CLAVE) se saltan."""
    from django.core.exceptions import ValidationError

    from apps.gastos.models import CentroCosto, ConceptoGasto, Gasto, Vehiculo

    centros = {c.codigo.upper(): c for c in CentroCosto.objects.filter(is_active=True).exclude(codigo=None)}
    conceptos = {normalizar(c.nombre): c for c in ConceptoGasto.objects.filter(is_active=True).select_related("grupo")}
    vehiculos = {normalizar(v.nombre): v for v in Vehiculo.objects.filter(is_active=True)}
    ya_importadas = set(Gasto.objects.exclude(clave_importacion=None).values_list("clave_importacion", flat=True))

    resultado = ResultadoImportacion()
    claves = set()
    for fila in filas:
        accion = normalizar(fila["accion"])
        resultado.por_accion[accion or "(vacía)"] = resultado.por_accion.get(accion or "(vacía)", 0) + 1
        etiqueta = f"Fila {fila['fila_revision']} ({_texto(fila['clave']) or 'sin clave'})"
        if accion not in ACCIONES:
            resultado.errores.append(f"{etiqueta}: ACCIÓN «{fila['accion']}» no válida (usa {', '.join(ACCIONES)}).")
            continue
        if accion != IMPORTAR:
            continue

        problemas = []
        clave = _texto(fila["clave"])
        if not clave:
            problemas.append("falta la CLAVE")
        elif clave in claves:
            problemas.append("CLAVE repetida; si repartiste un vale en varias filas, agrega /1, /2... a la clave")
        claves.add(clave)
        if clave in ya_importadas:
            resultado.ya_importadas += 1
            continue

        fecha = a_fecha(fila["fecha"])
        if fecha is None:
            problemas.append("FECHA vacía o no válida")
        importe = a_importe(fila["importe"])
        if importe is None:
            problemas.append("IMPORTE vacío o no mayor a cero")
        descripcion = _recortar(fila["descripcion"], 255)
        if not descripcion:
            problemas.append("DESCRIPCIÓN vacía")
        centro = centros.get(codigo_de_centro(fila["centro"]))
        if centro is None:
            problemas.append(f"CENTRO DE COSTO «{_texto(fila['centro'])}» no existe")
        concepto = conceptos.get(normalizar(fila["concepto"]))
        if concepto is None:
            problemas.append(f"CONCEPTO «{_texto(fila['concepto'])}» no existe")
        vehiculo = None
        if _texto(fila["vehiculo"]):
            vehiculo = vehiculos.get(normalizar(fila["vehiculo"]))
            if vehiculo is None:
                problemas.append(f"VEHÍCULO «{_texto(fila['vehiculo'])}» no existe")
        facturado_texto = normalizar(fila["facturado"])
        if facturado_texto not in ("SI", "NO"):
            problemas.append("FACTURADO debe ser SI o NO")
        facturado = facturado_texto == "SI"
        if problemas:
            resultado.errores.append(f"{etiqueta}: {'; '.join(problemas)}.")
            continue

        hoja, _, numero_fila = clave.partition("!")
        gasto = Gasto(
            centro_costo=centro,
            concepto_gasto=concepto,
            vehiculo=vehiculo,
            descripcion=descripcion,
            referencia=_recortar(fila["vale"], 100),
            condicion=Gasto.Condicion.PAGADO,
            fecha=fecha,
            importe=importe,
            facturado=facturado,
            referencia_factura=(_recortar(fila["referencia_factura"], 50) or FOLIO_PENDIENTE) if facturado else "",
            importado=True,
            turno_anterior=_recortar(fila["turno_anterior"], 20),
            factura_caja_anterior=_recortar(fila["factura_caja"], 30),
            clave_importacion=clave,
            observaciones=(
                f"Importado del Excel de pólizas: hoja {hoja}, fila {numero_fila}. "
                f"Sección: {_texto(fila['seccion']) or '—'}. Centro en el Excel: {_texto(fila['centro_original']) or '—'}. "
                f"Condición: {_texto(fila['condicion']) or '—'}."
            ),
        )
        try:
            gasto.full_clean(exclude=["comprobante"])
        except ValidationError as error:
            mensajes = [m for lista in error.message_dict.values() for m in lista]
            resultado.errores.append(f"{etiqueta}: {' '.join(mensajes)}")
            continue
        resultado.gastos.append(gasto)
    return resultado


def importar_revision(filas, aplicar=False):
    """Valida el archivo completo y, si `aplicar` y no hay ningún error,
    crea todos los gastos en una sola transacción (todo o nada)."""
    from django.db import transaction

    resultado = preparar_gastos(filas)
    if resultado.errores or not aplicar:
        return resultado
    with transaction.atomic():
        for gasto in resultado.gastos:
            gasto.save()
    resultado.creados = len(resultado.gastos)
    return resultado
