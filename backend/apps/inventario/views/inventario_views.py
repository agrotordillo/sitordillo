from decimal import Decimal

import openpyxl
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from django.views.generic import ListView
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.compras.models import OrdenCompra
from apps.core.filtros_producto import FiltrosProductoMixin
from apps.core.parametros import fecha, filtrar_por_id, id_valido
from apps.core.scoping import almacenes_visibles
from apps.inventario.models import Lote
from apps.products.busqueda import buscar_productos


class LoteListView(FiltrosProductoMixin, PermissionRequiredMixin, ListView):
    permission_required = "inventario.view_lote"
    model = Lote
    template_name = "inventario/lote_list.html"
    context_object_name = "lotes"
    extra_context = {"active_module": "warehouses"}
    filtro_producto_prefix = "producto__"
    filtro_incluir_almacen = True

    def get_queryset(self):
        queryset = (
            super()
            .get_queryset()
            .filter(cantidad_disponible__gt=0)
            .select_related("producto", "almacen", "orden_compra_detalle__orden_compra")
            .order_by("fecha_caducidad", "fecha_ingreso")
        )
        visibles = almacenes_visibles(self.request.user)
        if visibles is not None:
            queryset = queryset.filter(almacen__in=visibles)
        queryset = buscar_productos(queryset, self.request.GET.get("q", ""), "producto__")
        return filtrar_por_id(queryset, "orden_compra_detalle__orden_compra_id", self.request.GET.get("orden"))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["q"] = self.request.GET.get("q", "").strip()
        orden_id = id_valido(self.request.GET.get("orden"))
        if orden_id:
            context["orden"] = OrdenCompra.objects.filter(pk=orden_id).first()
        return context


class ExistenciaListView(FiltrosProductoMixin, PermissionRequiredMixin, ListView):
    permission_required = "inventario.view_lote"
    template_name = "inventario/existencia_list.html"
    context_object_name = "existencias"
    extra_context = {"active_module": "warehouses"}
    filtro_producto_prefix = "producto__"
    filtro_incluir_almacen = True

    def get_queryset(self):
        queryset = Lote.objects.filter(is_active=True, cantidad_disponible__gt=0)
        visibles = almacenes_visibles(self.request.user)
        if visibles is not None:
            queryset = queryset.filter(almacen__in=visibles)

        # Filtros de marca/línea/categoría/subcategoría/clase/sucursal deben
        # aplicarse ANTES del values()+annotate() de abajo: después de
        # agrupar, un .filter() ya actuaría como HAVING sobre lo agregado,
        # no como filtro de qué lotes entran al agrupamiento.
        queryset = self.aplicar_filtros_producto(queryset)
        queryset = buscar_productos(queryset, self.request.GET.get("q", ""), "producto__")

        # Fecha de ingreso del lote: cuándo entró esa mercancía, no cuándo
        # se registró en el sistema (created_at es solo auditoría).
        fecha_desde = fecha(self.request.GET.get("fecha_desde"))
        if fecha_desde:
            queryset = queryset.filter(fecha_ingreso__gte=fecha_desde)
        fecha_hasta = fecha(self.request.GET.get("fecha_hasta"))
        if fecha_hasta:
            queryset = queryset.filter(fecha_ingreso__lte=fecha_hasta)

        agrupado = (
            queryset.values("producto_id", "producto__nombre", "producto__sku", "almacen_id", "almacen__nombre")
            .annotate(total=Sum("cantidad_disponible"))
            .order_by("producto__nombre", "producto_id")
        )
        self.almacenes, filas = pivotar_por_almacen(agrupado)
        return filas

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["almacenes"] = self.almacenes
        context["q"] = self.request.GET.get("q", "").strip()
        context["fecha_desde"] = self.request.GET.get("fecha_desde", "")
        context["fecha_hasta"] = self.request.GET.get("fecha_hasta", "")
        context["querystring"] = self.request.GET.urlencode()
        return context


def pivotar_por_almacen(agrupado):
    """De una fila por (producto, almacén) a una fila por producto con una
    columna por almacén, como se leía en el sistema anterior. Las columnas
    son solo los almacenes con existencia en el resultado, por nombre; la
    celda queda en None donde ese producto no tiene existencia."""
    agrupado = list(agrupado)
    almacenes = sorted({(fila["almacen__nombre"], fila["almacen_id"]) for fila in agrupado})
    posicion = {almacen_id: indice for indice, (_, almacen_id) in enumerate(almacenes)}

    filas = {}
    for registro in agrupado:
        fila = filas.setdefault(registro["producto_id"], {
            "producto__nombre": registro["producto__nombre"],
            "producto__sku": registro["producto__sku"],
            "cantidades": [None] * len(almacenes),
            "total": Decimal("0"),
        })
        fila["cantidades"][posicion[registro["almacen_id"]]] = registro["total"]
        fila["total"] += registro["total"]
    return [nombre for nombre, _ in almacenes], list(filas.values())


class ExistenciaExcelView(ExistenciaListView):
    """Descarga en Excel la misma tabla de Existencias que se ve en
    pantalla, con los mismos filtros."""

    def render_to_response(self, context, **response_kwargs):
        almacenes = context["almacenes"]
        workbook = openpyxl.Workbook()
        hoja = workbook.active
        hoja.title = "Existencias"
        hoja.append(["SKU", "Producto", *almacenes, "Total"])
        for celda in hoja[1]:
            celda.font = Font(bold=True)

        for fila in context["existencias"]:
            hoja.append([
                fila["producto__sku"],
                fila["producto__nombre"],
                *(float(cantidad) if cantidad is not None else None for cantidad in fila["cantidades"]),
                float(fila["total"]),
            ])

        hoja.column_dimensions["A"].width = 14
        hoja.column_dimensions["B"].width = 40
        for indice in range(3, len(almacenes) + 4):
            hoja.column_dimensions[get_column_letter(indice)].width = 16
            for (celda,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celda.number_format = "#,##0.00"
        hoja.freeze_panes = "C2"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        fecha_hoy = timezone.localdate().isoformat()
        response["Content-Disposition"] = f'attachment; filename="existencias_{fecha_hoy}.xlsx"'
        workbook.save(response)
        return response
