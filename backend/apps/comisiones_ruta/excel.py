"""Descarga a Excel del reporte de comisión por colaborador."""
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = "Arial"
MONEDA = '$#,##0.00;-$#,##0.00;"-"'
PORCENTAJE = '0.00"%"'
CANTIDAD = '#,##0.##;-#,##0.##;"-"'
RELLENO = PatternFill("solid", fgColor="334155")
TOTAL = PatternFill("solid", fgColor="E2E8F0")
FORMATOS = {None: "texto", MONEDA: "moneda", PORCENTAJE: "porcentaje", CANTIDAD: "cantidad"}


def _nombre(usuario):
    return usuario.get_full_name() or usuario.username


def _titulo_de_hoja(nombre, usados):
    """Nombre de hoja válido para Excel (31 caracteres, sin []:*?/\\) y único."""
    base = re.sub(r"[\[\]:*?/\\]", " ", nombre).strip()[:31] or "Colaborador"
    titulo, n = base, 2
    while titulo.lower() in usados:
        sufijo = f" ({n})"
        titulo, n = base[: 31 - len(sufijo)] + sufijo, n + 1
    usados.add(titulo.lower())
    return titulo


def _encabezado(hoja, fila, columnas):
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(fila, i, titulo)
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = RELLENO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = max(hoja.column_dimensions[get_column_letter(i)].width or 0, ancho)


def _estilos(libro, formatos, total=False):
    """El nombre del estilo de celda (NamedStyle) de cada columna, registrado
    en el libro la primera vez. Asignar Font y number_format celda por celda
    obliga a openpyxl a crear y deduplicar un objeto por celda: con decenas
    de miles de filas de detalle era lo que más tardaba la descarga."""
    nombres = []
    for formato in formatos:
        nombre = f"Comisiones {FORMATOS[formato]}" + (" total" if total else "")
        if nombre not in libro.style_names:
            estilo = NamedStyle(nombre, font=Font(name=FUENTE, size=10, bold=total), number_format=formato or "General")
            if total:
                estilo.fill = TOTAL
            libro.add_named_style(estilo)
        nombres.append(nombre)
    return nombres


def _escribir(hoja, fila, valores, estilos):
    for i, (valor, estilo) in enumerate(zip(valores, estilos), start=1):
        hoja.cell(fila, i, valor).style = estilo


def libro_comisiones_ruta(contexto):
    """Un resumen (colaborador, vendido, devuelto, comisión y lo que se
    descontará después) y una hoja por colaborador: sus reglas con la base y
    la comisión del periodo, como los bloques del Excel con el que se pagan;
    las devoluciones que se descuentan en el periodo y las posteriores con
    el periodo en que se descuentan; y el detalle por venta. La comisión de
    cada regla es la suma de la de cada venta y devolución (redondeada
    renglón por renglón, igual que en pantalla); los totales van como
    fórmulas."""
    desde, hasta = contexto["periodo"]
    filtros = f"Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"
    if contexto.get("almacen"):
        filtros += f" · Sucursal: {contexto['almacen'].nombre}"
    filtros += f" · Generado el {timezone.localtime():%d/%m/%Y %H:%M}"

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen["A1"] = "Comisiones por colaborador"
    resumen["A1"].font = Font(name=FUENTE, size=14, bold=True)
    resumen["A2"] = filtros
    resumen["A2"].font = Font(name=FUENTE, size=9, italic=True)
    _encabezado(resumen, 4, [
        ("Colaborador", 40), ("Plaza", 18), ("Vendido (sin IVA)", 18), ("Devuelto (sin IVA)", 18), ("Comisión", 16),
        ("Devoluciones posteriores: comisión a descontar después", 22),
    ])
    usados = {"resumen"}
    fila_resumen = 5
    formatos = [None, None, MONEDA, MONEDA, MONEDA, MONEDA]
    estilos = _estilos(libro, formatos)
    for grupo in contexto["grupos"]:
        _escribir(
            resumen, fila_resumen,
            [_nombre(grupo["empleado"]), grupo["plaza"].nombre if grupo["plaza"] else "",
             float(grupo["total_importe"]), -float(grupo["total_devuelto"]), float(grupo["total_comision"]),
             float(grupo["total_por_descontar"])],
            estilos,
        )
        _hoja_del_colaborador(libro, grupo, filtros, usados)
        fila_resumen += 1
    ultima = fila_resumen - 1
    _escribir(
        resumen, fila_resumen,
        ["Total", "", *(f"=SUM({c}5:{c}{ultima})" if ultima >= 5 else 0 for c in "CDEF")],
        _estilos(libro, formatos, total=True),
    )
    notas = []
    if contexto.get("sin_lista_precio"):
        notas.append(
            f"{contexto['sin_lista_precio']} línea(s) de venta con comisión configurada no se incluyeron por no "
            "tener registrada su lista de precio."
        )
    notas.append(
        "Una devolución se descuenta en el periodo en que se registra, con el % con que se pagó la venta. Las "
        "devoluciones posteriores no cambian la comisión de este periodo: se descuentan en el periodo que indica "
        "la hoja de cada colaborador."
    )
    for i, nota in enumerate(notas):
        resumen.cell(fila_resumen + 2 + i, 1, nota).font = Font(name=FUENTE, size=9, italic=True)
    resumen.freeze_panes = "A5"
    return libro


def _periodo(periodo):
    return f"{periodo[0]:%d/%m/%Y} al {periodo[1]:%d/%m/%Y}"


def _tabla(hoja, fila, titulo, columnas, formatos, renglones, sumar):
    """Escribe `titulo`, el encabezado y los renglones a partir de `fila`, y
    un renglón de total con fórmulas en las columnas `sumar` (letras).
    Devuelve el número de fila siguiente al total."""
    libro = hoja.parent
    hoja.cell(fila, 1, titulo).font = Font(name=FUENTE, bold=True)
    _encabezado(hoja, fila + 1, columnas)
    estilos = _estilos(libro, formatos)
    primera = fila + 2
    actual = primera
    for valores in renglones:
        _escribir(hoja, actual, valores, estilos)
        actual += 1
    total = ["Total"] + [""] * (len(columnas) - 1)
    for letra in sumar:
        total[ord(letra) - ord("A")] = f"=SUM({letra}{primera}:{letra}{actual - 1})"
    _escribir(hoja, actual, total, _estilos(libro, formatos, total=True))
    return actual + 1


def _hoja_del_colaborador(libro, grupo, filtros, usados):
    hoja = libro.create_sheet(_titulo_de_hoja(_nombre(grupo["empleado"]), usados))
    hoja["A1"] = _nombre(grupo["empleado"]) + (f" · {grupo['plaza'].nombre}" if grupo["plaza"] else "")
    hoja["A1"].font = Font(name=FUENTE, size=13, bold=True)
    hoja["A2"] = filtros
    hoja["A2"].font = Font(name=FUENTE, size=9, italic=True)

    # Reglas
    _encabezado(hoja, 4, [
        ("Regla", 60), ("Vendido PUBLICO / M. MAYOREO", 18), ("%", 9), ("Vendido otras listas", 16),
        ("% otras", 9), ("Devuelto", 14), ("Comisión", 14),
    ])
    fila = 5
    formatos = [None, MONEDA, PORCENTAJE, MONEDA, PORCENTAJE, MONEDA, MONEDA]
    estilos = _estilos(libro, formatos)
    for totales in grupo["reglas"]:
        regla = totales["regla"]
        _escribir(
            hoja, fila,
            [regla.descripcion, float(totales["vendido_variable"]), float(regla.porcentaje),
             float(totales["vendido_otras"]), float(regla.porcentaje_otras_listas), -float(totales["devuelto"]),
             float(totales["comision"])],
            estilos,
        )
        fila += 1
    ultima = fila - 1
    _escribir(
        hoja, fila,
        ["Total", f"=SUM(B5:B{ultima})", "", f"=SUM(D5:D{ultima})", "", f"=SUM(F5:F{ultima})", f"=SUM(G5:G{ultima})"],
        _estilos(libro, [None, MONEDA, None, MONEDA, None, MONEDA, MONEDA], total=True),
    )
    fila += 3

    zona = timezone.get_current_timezone()
    descripciones = {t["regla"].pk: t["regla"].descripcion for t in grupo["reglas"]}

    # Devoluciones: las que se descuentan aquí y las posteriores.
    columnas_devolucion = [
        ("Devolución", 16), ("Fecha devolución", 13), ("Venta", 16), ("Fecha venta", 17), (None, 24),
        ("Sucursal", 18), ("Vendió", 28), ("Producto", 40), ("SKU", 14), ("Cantidad", 11),
        ("Importe (sin IVA)", 15), ("%", 9), ("Comisión", 13), ("Regla", 60),
    ]
    formatos_devolucion = [None] * 9 + [CANTIDAD, MONEDA, PORCENTAJE, MONEDA, None]
    for clave, titulo, periodo, encabezado_periodo in [
        ("devoluciones", "Devoluciones descontadas en este periodo", "periodo_venta", "Venta pagada en el periodo"),
        ("devoluciones_posteriores",
         "Devoluciones posteriores de ventas de este periodo (no cambian esta comisión; se descuentan en el periodo indicado)",
         "periodo_descuento", "Se descuenta en el periodo"),
    ]:
        if not grupo[clave]:
            continue
        columnas = [(nombre or encabezado_periodo, ancho) for nombre, ancho in columnas_devolucion]
        fila = _tabla(hoja, fila, titulo, columnas, formatos_devolucion, [
            [f["devolucion"].folio, f"{f['devolucion'].fecha:%d/%m/%Y}", f["venta"].folio,
             f["venta"].fecha_venta.astimezone(zona).strftime("%d/%m/%Y %H:%M"), _periodo(f[periodo]),
             f["venta"].almacen.nombre, _nombre(f["vendedor"]) if f["vendedor"] else "", f["producto"].nombre,
             f["producto"].sku, float(f["cantidad"]), float(f["importe"]), float(f["porcentaje"]),
             float(f["comision"]), descripciones[f["regla"].pk]]
            for f in grupo[clave]
        ], sumar="KM") + 2

    # Detalle por venta
    if grupo["filas"]:
        _tabla(hoja, fila, "Detalle por venta", [
            ("Fecha", 17), ("Folio", 16), ("Sucursal", 18), ("Vendió", 28), ("Producto", 40), ("SKU", 14),
            ("Lista", 14), ("Cantidad", 11), ("Importe (sin IVA)", 15), ("%", 9), ("Comisión", 13), ("Regla", 60),
        ], [None, None, None, None, None, None, None, CANTIDAD, MONEDA, PORCENTAJE, MONEDA, None], (
            [f["venta"].fecha_venta.astimezone(zona).strftime("%d/%m/%Y %H:%M"), f["venta"].folio,
             f["venta"].almacen.nombre, _nombre(f["vendedor"]) if f["vendedor"] else "", f["producto"].nombre,
             f["producto"].sku, f["lista_precio"].nombre if f["lista_precio"] else "", float(f["cantidad"]),
             float(f["importe"]), float(f["porcentaje"]), float(f["comision"]), descripciones[f["regla"].pk]]
            for f in grupo["filas"]
        ), sumar="IK")
    hoja.freeze_panes = "A5"
