import shutil
import tempfile
from datetime import date, datetime
from decimal import Decimal
from io import StringIO

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.clientes.models import Cliente
from apps.comisiones_ruta.models import ComisionColaboradorLinea
from apps.comisiones_ruta.views.reporte_views import periodo_de_pago
from apps.products.models import Almacen, Categoria, Clase, Linea, ListaPrecio, Producto
from apps.ventas.models import DevolucionCliente, DevolucionClienteDetalle, Venta, VentaDetalle

User = get_user_model()


class ImportComisionesExcelTests(TestCase):
    """Carga de los % por colaborador y línea desde el Excel de comisiones
    (import_comisiones_excel): la línea sale de los productos del bloque, no de
    su título, y lo que no es una línea se deja fuera."""

    @classmethod
    def setUpTestData(cls):
        cls.alimentos = Linea.objects.create(nombre="ALIMENTOS")
        cls.bombas = Linea.objects.create(nombre="REFACCIONES PARA BOMBAS")
        cls.semillas = Linea.objects.create(nombre="SEMILLAS Y FERTILIZANTES")
        Producto.objects.create(nombre="CROQUETA 20KG", sku="CRQ", linea=cls.alimentos)
        Producto.objects.create(nombre="SELLO MECANICO", sku="SEL", linea=cls.bombas)
        Producto.objects.create(nombre="TRIPLE 17 POR KG", sku="T17", linea=cls.semillas)
        cls.guillermo = User.objects.create(username="guillermo", is_active=False)

    def setUp(self):
        directorio = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, directorio)
        self.archivo = f"{directorio}/p36.xlsx"
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = "GUILLERMO"
        hoja.append(["BS", "GUILLERMO VALENZUELA"])
        fila = 2
        # (título, producto, % de Precio 1-2): el tercero es un bloque de
        # semillas que quedó con el título del anterior, como en el Excel real.
        for titulo, producto, pct in [
            ("ALIMENTOS", "CROQUETA 20KG", 3),
            ("REFACCIONES PARA BOMBAS", "SELLO MECANICO", 10),
            ("REFACCIONES PARA BOMBAS", "TRIPLE 17 POR KG", 5),
            ("SOBRES", "CROQUETA 20KG", 10),
        ]:
            hoja.cell(fila, 2, titulo)
            hoja.cell(fila + 1, 2, "Registros")
            hoja.cell(fila + 2, 1, producto)
            hoja.cell(fila + 3, 1, "Total")
            hoja.cell(fila + 5, 5, f"=+E{fila + 4}*{pct}%")
            hoja.cell(fila + 5, 7, f"=+G{fila + 4}*1%")
            fila += 7
        libro.save(self.archivo)

    def test_carga_por_linea_segun_los_productos_del_bloque(self):
        salida = StringIO()
        call_command("import_comisiones_excel", archivo=self.archivo, stdout=salida)

        cargados = dict(
            ComisionColaboradorLinea.objects.filter(colaborador=self.guillermo).values_list("linea__nombre", "porcentaje")
        )
        self.assertEqual(cargados, {
            "ALIMENTOS": Decimal("3.00"),
            "REFACCIONES PARA BOMBAS": Decimal("10.00"),
            "SEMILLAS Y FERTILIZANTES": Decimal("5.00"),
        })
        self.assertIn("trae productos de SEMILLAS Y FERTILIZANTES", salida.getvalue())
        # Sin la categoría SOBRES en el catálogo, el bloque no se puede cargar.
        self.assertIn("'SOBRES' no es línea, categoría ni grupo conocido", salida.getvalue())

    def test_dry_run_no_guarda(self):
        call_command("import_comisiones_excel", archivo=self.archivo, dry_run=True, stdout=StringIO())

        self.assertFalse(ComisionColaboradorLinea.objects.exists())

    def _agregar_hojas(self, **hojas):
        """Agrega al Excel de prueba hojas {nombre: {celda: valor}}."""
        with open(self.archivo, "rb") as archivo:
            libro = openpyxl.load_workbook(archivo)
        for nombre, celdas in hojas.items():
            hoja = libro.create_sheet(nombre)
            for celda, valor in celdas.items():
                hoja[celda] = valor
        libro.save(self.archivo)

    def test_hojas_de_total_categorias_y_farmacia_responsiva_de_la_sucursal(self):
        Categoria.objects.create(nombre="SOBRES")
        medicos = Clase.objects.create(nombre="MEDICOS")
        lerdo = Almacen.objects.create(nombre="LERDO", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        frank = User.objects.create(username="frank", is_active=False)
        alberto = User.objects.create(username="alberto", is_active=False)
        self._agregar_hojas(
            FRANKLIN={"B14": "TOTAL A PERCIBIR", "D14": "=+D11*0.5%"},
            **{"MVZ ALBERTO": {"C3": "FARMACIA RESPONSIVA", "C4": "Registros", "F7": "=+F6*4%", "H7": "=+H6*1%"}},
        )

        call_command("import_comisiones_excel", archivo=self.archivo, stdout=StringIO())

        # FRANKLIN: 0.5% de toda su venta, incluido lo vendido con precio especial.
        regla = ComisionColaboradorLinea.objects.get(colaborador=frank)
        self.assertEqual(
            (regla.linea, regla.porcentaje, regla.porcentaje_otras_listas, regla.comisiona_precio_especial),
            (None, Decimal("0.50"), Decimal("0.50"), True),
        )
        regla = ComisionColaboradorLinea.objects.get(colaborador=alberto)
        self.assertEqual(
            (regla.alcance, regla.clase, regla.almacen, regla.porcentaje),
            (ComisionColaboradorLinea.Alcance.SUCURSAL, medicos, lerdo, 4),
        )
        # El bloque SOBRES de GUILLERMO ahora es una regla de categoría.
        self.assertTrue(ComisionColaboradorLinea.objects.filter(
            colaborador=self.guillermo, categoria__nombre="SOBRES", porcentaje=10,
        ).exists())

    def test_jhony_cobra_distinto_segun_la_sucursal_y_se_avisa_el_ajuste_a_mano(self):
        for numero, nombre in enumerate(["IQUINUAPA", "TECOLUTILLA", "LERDO"], start=1):
            Almacen.objects.create(nombre=nombre, tipo=Almacen.Tipo.SUCURSAL, numero=numero)
        jhony = User.objects.create(username="jhony", is_active=False)
        self._agregar_hojas(JHONY={
            "B2": "COMISIONES DEL 15 AL 21 DE OCTUBRE 2026 P43",
            "D13": "=+D12*1%+782.63",
            "B14": "REFACCIONES PARA BOMBAS", "B15": "Registros", "A16": "SELLO MECANICO", "E19": "=+E18*10%",
            "B21": "ALIMENTOS", "B22": "Registros", "A23": "CROQUETA 20KG", "E26": "=+E25*1%", "G26": "=+G25*1%",
        })
        salida = StringIO()

        call_command("import_comisiones_excel", archivo=self.archivo, stdout=salida)

        Alcance = ComisionColaboradorLinea.Alcance
        reglas = {
            (r.almacen.nombre, r.linea.nombre if r.linea else None): (
                r.alcance, r.porcentaje, r.comisiona_precio_especial, r.dias_semana, r.vigente_desde,
            )
            for r in ComisionColaboradorLinea.objects.filter(colaborador=jhony)
        }
        self.assertEqual(reglas, {
            # Lo que vende en IQUINUAPA: 1% de todo.
            ("IQUINUAPA", None): (Alcance.PROPIAS_EN_SUCURSAL, Decimal("1.00"), True, "0123456", None),
            # En LERDO cubrió el esquema de MIGUEL el domingo de esa semana: toda la venta de la clase.
            ("LERDO", "REFACCIONES PARA BOMBAS"): (Alcance.SUCURSAL, Decimal("10.00"), False, "6", date(2026, 10, 15)),
            ("TECOLUTILLA", "ALIMENTOS"): (Alcance.PROPIAS_EN_SUCURSAL, Decimal("1.00"), False, "0123456", None),
        })
        self.assertIn("suma a mano +782.63", salida.getvalue())

    def test_esquema_por_clase_con_descanso_y_cobertura_con_las_fechas_de_su_hoja(self):
        lerdo = Almacen.objects.create(nombre="LERDO", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        jose_manuel = Clase.objects.create(nombre="JOSE MANUEL")
        Clase.objects.create(nombre="MIGUEL RAMIREZ")
        luis = User.objects.create(username="luisg", is_active=False)
        bloque = {"C3": "REFACCIONES PARA BOMBAS", "C4": "Registros", "B5": "SELLO MECANICO", "F8": "=+F7*5%",
                  "H8": "=+H7*1%"}
        self._agregar_hojas(**{
            "LUIS GDO": bloque,
            "JUAN TARACENA": {"B2": "COMISIONES DEL 15 AL 21 DE OCTUBRE 2026 P43", **{
                "C3": "REFACCIONES PARA BOMBAS", "C4": "Registros", "B5": "SELLO MECANICO", "F8": "=+F7*10%"}},
        })
        salida = StringIO()

        call_command("import_comisiones_excel", archivo=self.archivo, stdout=salida)

        # LUIS (descansa jueves): la línea sobre toda la venta de LERDO de su clase.
        regla = ComisionColaboradorLinea.objects.get(colaborador=luis)
        self.assertEqual(
            (regla.alcance, regla.almacen, regla.clase, regla.linea, regla.porcentaje, regla.dias_semana),
            (ComisionColaboradorLinea.Alcance.SUCURSAL, lerdo, jose_manuel, self.bombas, Decimal("5.00"), "012456"),
        )
        # TARACENA no vende en scv_web: se le crea un usuario inactivo, y su regla
        # es la del esquema de MIGUEL solo en las fechas de su hoja, sin domingo.
        taracena = User.objects.get(username="jtaracena")
        regla = ComisionColaboradorLinea.objects.get(colaborador=taracena)
        self.assertEqual(
            (taracena.is_active, regla.alcance, regla.almacen, regla.clase.nombre, regla.linea, regla.dias_semana,
             regla.vigente_desde, regla.vigente_hasta),
            (False, ComisionColaboradorLinea.Alcance.SUCURSAL, lerdo, "MIGUEL RAMIREZ", self.bombas, "012345",
             date(2026, 10, 15), date(2026, 10, 21)),
        )
        self.assertIn("se creó el usuario inactivo jtaracena", salida.getvalue())

    def test_lee_el_periodo_del_encabezado_de_la_hoja(self):
        from apps.comisiones_ruta.management.commands.import_comisiones_excel import periodo_de_la_hoja

        casos = {
            "COMISIONES DEL 15 AL 21 DE OCTUBRE 2026 P43": (date(2026, 10, 15), date(2026, 10, 21)),
            "COMISIONES DEL 15 AL 16 y 18 AL 21 DE OCTUBRE 2026 P43": (date(2026, 10, 15), date(2026, 10, 21)),
            "COMISIONES DEL 27 DE AGOSTO AL 02 DE SEPRIEMBRE 2026 P36": (date(2026, 8, 27), date(2026, 9, 2)),
            "COMISIONES DEL 18 DE OCTUBRE 2026 P43 2026": (date(2026, 10, 18), date(2026, 10, 18)),
        }
        for encabezado, esperado in casos.items():
            hoja = openpyxl.Workbook().active
            hoja["B2"] = encabezado
            self.assertEqual(periodo_de_la_hoja(hoja), esperado, encabezado)

    def test_recargar_reemplaza_lo_de_un_excel_y_respeta_lo_capturado_a_mano(self):
        manual = ComisionColaboradorLinea.objects.create(
            colaborador=self.guillermo, linea=self.semillas, porcentaje=7, observaciones="Capturada a mano",
            alcance=ComisionColaboradorLinea.Alcance.PROPIAS_EN_SUCURSAL,
            almacen=Almacen.objects.create(nombre="LERDO", tipo=Almacen.Tipo.SUCURSAL, numero=1),
        )
        call_command("import_comisiones_excel", archivo=self.archivo, stdout=StringIO())
        call_command("import_comisiones_excel", archivo=self.archivo, stdout=StringIO())

        self.assertEqual(ComisionColaboradorLinea.objects.filter(observaciones__startswith="Cargado de ").count(), 3)
        self.assertTrue(ComisionColaboradorLinea.objects.filter(pk=manual.pk).exists())


def _en(dia):
    return timezone.make_aware(datetime(2026, 9, dia, 12, 0))


class ReporteComisionesRutaTests(TestCase):
    """Reporte de comisión por colaborador con reglas por línea, categoría,
    clase y sobre la venta de una sucursal (ComisionColaboradorLinea)."""

    @classmethod
    def setUpTestData(cls):
        cls.lerdo = Almacen.objects.create(nombre="LERDO", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        cls.publico = ListaPrecio.objects.get_or_create(nombre="PUBLICO")[0]
        cls.mayoreo = ListaPrecio.objects.get_or_create(nombre="MAYOREO")[0]
        cls.cliente = Cliente.objects.create(nombre="Cliente")
        farmacia = Linea.objects.create(nombre="FARMACIA")
        alimentos = Linea.objects.create(nombre="ALIMENTOS")
        sobres = Categoria.objects.create(nombre="SOBRES")
        medicos = Clase.objects.create(nombre="MEDICOS")
        sin_iva = {"tipo_iva": Producto.TipoIVA.GRAVADO, "tasa_iva": Decimal("0.00")}
        cls.antibiotico = Producto.objects.create(nombre="Antibiótico", sku="ANT", linea=farmacia, clase=medicos, **sin_iva)
        cls.sobre = Producto.objects.create(nombre="Sobre", sku="SOB", linea=farmacia, categoria=sobres, **sin_iva)
        cls.croqueta = Producto.objects.create(nombre="Croqueta", sku="CRQ", linea=alimentos, **sin_iva)
        cls.luis = User.objects.create(username="luisg", first_name="LUIS", is_active=False)
        cls.alberto = User.objects.create(username="alberto", first_name="ALBERTO", is_active=False)
        cls.marlit = User.objects.create(username="mgonzalez", first_name="MARLIT", is_active=False)
        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")

        cls.iquinuapa = Almacen.objects.create(nombre="IQUINUAPA", tipo=Almacen.Tipo.SUCURSAL, numero=2)
        cls.jhony = User.objects.create(username="jhony", first_name="JHONY", is_active=False)
        Alcance = ComisionColaboradorLinea.Alcance

        regla = ComisionColaboradorLinea.objects.create
        regla(colaborador=cls.luis, linea=farmacia, porcentaje=5)
        regla(colaborador=cls.luis, categoria=sobres, porcentaje=10)
        regla(colaborador=cls.alberto, alcance=Alcance.SUCURSAL, almacen=cls.lerdo, clase=medicos, porcentaje=4)
        regla(colaborador=cls.marlit, linea=alimentos, porcentaje=1, porcentaje_otras_listas=0)
        # JHONY: 1% de todo lo que vende en IQUINUAPA (incluido precio especial);
        # en LERDO, solo ALIMENTOS al 3%.
        regla(colaborador=cls.jhony, alcance=Alcance.PROPIAS_EN_SUCURSAL, almacen=cls.iquinuapa,
              porcentaje=1, porcentaje_otras_listas=1, comisiona_precio_especial=True)
        regla(colaborador=cls.jhony, alcance=Alcance.PROPIAS_EN_SUCURSAL, almacen=cls.lerdo, linea=alimentos,
              porcentaje=3)

        def venta(vendedor, producto, importe, lista, almacen=None, especial=False):
            v = Venta.objects.create(
                cliente=cls.cliente, almacen=almacen or cls.lerdo, fecha_venta=_en(2), created_by=vendedor,
            )
            VentaDetalle.objects.create(
                venta=v, producto=producto, cantidad=1, precio_unitario=importe, lista_precio=lista,
                precio_especial=especial,
            )

        venta(cls.luis, cls.antibiotico, Decimal("1000.00"), cls.publico)
        venta(cls.luis, cls.sobre, Decimal("200.00"), cls.publico)
        venta(cls.marlit, cls.croqueta, Decimal("500.00"), cls.publico)
        venta(cls.marlit, cls.croqueta, Decimal("800.00"), cls.mayoreo)
        # Precio especial de sucursal: no comisiona en la regla por línea de MARLIT.
        venta(cls.marlit, cls.croqueta, Decimal("700.00"), cls.publico, especial=True)
        venta(cls.jhony, cls.croqueta, Decimal("300.00"), cls.publico, almacen=cls.iquinuapa, especial=True)
        venta(cls.jhony, cls.sobre, Decimal("100.00"), cls.publico, almacen=cls.iquinuapa)
        venta(cls.jhony, cls.croqueta, Decimal("100.00"), cls.publico)
        venta(cls.jhony, cls.sobre, Decimal("100.00"), cls.publico)  # LERDO, sin regla de FARMACIA

    def setUp(self):
        self.client.force_login(self.admin)

    def _comisiones(self, **filtros):
        respuesta = self.client.get(
            reverse("comisiones_ruta:reporte"), {"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30", **filtros}
        )
        self.assertEqual(respuesta.status_code, 200)
        return {g["empleado"].username: g["total_comision"] for g in respuesta.context["grupos"]}, respuesta

    def test_cada_colaborador_cobra_segun_su_regla_mas_especifica(self):
        comisiones, _ = self._comisiones()

        self.assertEqual(comisiones, {
            # FARMACIA 5% sobre el antibiótico + SOBRES 10% (más específica) sobre el sobre.
            "luisg": Decimal("70.00"),
            # 4% del antibiótico de clase MEDICOS que vendió otro en LERDO.
            "alberto": Decimal("40.00"),
            # 1% de la croqueta en PUBLICO; la de MAYOREO va al 0% y la de precio especial no cuenta.
            "mgonzalez": Decimal("5.00"),
            # IQUINUAPA 1% de 300 (precio especial incluido) + 100; LERDO 3% de 100 en ALIMENTOS.
            "jhony": Decimal("7.00"),
        })

    def test_la_regla_no_aplica_en_su_dia_de_descanso_ni_fuera_de_su_vigencia(self):
        regla = ComisionColaboradorLinea.objects.get(colaborador=self.alberto)
        # Las ventas son del miércoles 2 de septiembre de 2026.
        regla.dias_semana = "013456"  # descansa los miércoles
        regla.save()
        self.assertNotIn("alberto", self._comisiones()[0])

        regla.dias_semana = "0123456"
        regla.vigente_desde, regla.vigente_hasta = date(2026, 9, 3), date(2026, 9, 5)  # una cobertura posterior
        regla.save()
        self.assertNotIn("alberto", self._comisiones()[0])

        regla.vigente_desde = date(2026, 9, 1)
        regla.save()
        self.assertEqual(self._comisiones()[0]["alberto"], Decimal("40.00"))

    def test_se_descarga_en_excel_con_una_hoja_por_colaborador_y_su_detalle(self):
        import io

        respuesta = self.client.get(
            reverse("comisiones_ruta:reporte-excel"), {"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30"}
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('filename="comisiones_20260901_20260930.xlsx"', respuesta["Content-Disposition"])
        libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ["Resumen", "ALBERTO", "JHONY", "LUIS", "MARLIT"])
        resumen = libro["Resumen"]
        self.assertEqual(
            [(resumen.cell(f, 1).value, resumen.cell(f, 5).value) for f in range(5, 9)],
            [("ALBERTO", 40), ("JHONY", 7), ("LUIS", 70), ("MARLIT", 5)],
        )
        self.assertEqual((resumen["A9"].value, resumen["E9"].value), ("Total", "=SUM(E5:E8)"))
        # La hoja de LUIS: sus 2 reglas, su total y el detalle de sus 2 ventas.
        luis = libro["LUIS"]
        self.assertEqual(luis["G7"].value, "=SUM(G5:G6)")
        self.assertEqual(luis["A10"].value, "Detalle por venta")
        folios = [luis.cell(f, 2).value for f in range(12, 14)]
        self.assertTrue(all(folios))
        self.assertEqual(luis.cell(14, 11).value, "=SUM(K12:K13)")

    def test_el_detalle_por_venta_solo_sale_al_filtrar_un_colaborador(self):
        _, respuesta = self._comisiones()
        self.assertTrue(all(not g["filas"] for g in respuesta.context["grupos"]))

        _, respuesta = self._comisiones(empleado=self.luis.pk)
        [grupo] = respuesta.context["grupos"]
        self.assertEqual(sorted(f["comision"] for f in grupo["filas"]), [Decimal("20.00"), Decimal("50.00")])

    # --- Devoluciones ---------------------------------------------------------

    def _devolver(self, vendedor, producto, dia):
        """Devuelve la pieza que `vendedor` vendió de `producto` el día `dia`
        de septiembre de 2026."""
        detalle = VentaDetalle.objects.get(venta__created_by=vendedor, producto=producto)
        devolucion = DevolucionCliente.objects.create(venta=detalle.venta, fecha=date(2026, 9, dia))
        DevolucionClienteDetalle.objects.create(devolucion=devolucion, venta_detalle=detalle, cantidad=1)
        return devolucion

    def _grupo(self, respuesta, colaborador):
        return next(g for g in respuesta.context["grupos"] if g["empleado"] == colaborador)

    def test_el_periodo_de_pago_va_de_jueves_a_miercoles(self):
        self.assertEqual(periodo_de_pago(date(2026, 10, 28)), (date(2026, 10, 22), date(2026, 10, 28)))
        self.assertEqual(periodo_de_pago(date(2026, 10, 29)), (date(2026, 10, 29), date(2026, 11, 4)))

    def test_la_devolucion_de_una_venta_ya_pagada_se_descuenta_en_el_periodo_en_que_se_registra(self):
        # La venta es del miércoles 2/9 (periodo de pago del 27/8 al 2/9) y
        # la devuelven el jueves 10/9, con ese periodo ya pagado.
        devolucion = self._devolver(self.luis, self.antibiotico, 10)

        pagado, respuesta = self._comisiones(fecha_desde="2026-08-27", fecha_hasta="2026-09-02")
        # El periodo pagado no cambia, pero avisa dónde se descontará.
        self.assertEqual((pagado["luisg"], pagado["alberto"]), (Decimal("70.00"), Decimal("40.00")))
        [posterior] = self._grupo(respuesta, self.luis)["devoluciones_posteriores"]
        self.assertEqual(
            (posterior["devolucion"], posterior["comision"], posterior["periodo_descuento"]),
            (devolucion, Decimal("-50.00"), (date(2026, 9, 10), date(2026, 9, 16))),
        )
        self.assertEqual(respuesta.context["total_por_descontar"], Decimal("-90.00"))
        self.assertContains(respuesta, "Devoluciones posteriores de ventas de este periodo")

        siguiente, respuesta = self._comisiones(fecha_desde="2026-09-10", fecha_hasta="2026-09-16")
        # Se retira lo que se pagó: a LUIS su 5% de FARMACIA y a ALBERTO su 4%
        # de MEDICOS en LERDO, aunque no hayan vendido nada en este periodo.
        self.assertEqual(siguiente, {"luisg": Decimal("-50.00"), "alberto": Decimal("-40.00")})
        luis = self._grupo(respuesta, self.luis)
        [descontada] = luis["devoluciones"]
        self.assertEqual(
            (descontada["devolucion"], descontada["importe"], descontada["periodo_venta"]),
            (devolucion, Decimal("-1000.00"), (date(2026, 8, 27), date(2026, 9, 2))),
        )
        self.assertEqual((luis["total_importe"], luis["total_devuelto"]), (Decimal("0.00"), Decimal("1000.00")))
        self.assertContains(respuesta, "Devoluciones descontadas en este periodo")

    def test_la_devolucion_en_el_mismo_periodo_se_descuenta_ahi_mismo(self):
        self._devolver(self.luis, self.antibiotico, 2)

        comisiones, respuesta = self._comisiones()

        self.assertEqual((comisiones["luisg"], comisiones["alberto"]), (Decimal("20.00"), Decimal("0.00")))
        self.assertFalse(any(g["devoluciones_posteriores"] for g in respuesta.context["grupos"]))

    def test_la_devolucion_retira_con_la_regla_del_dia_de_la_venta_aunque_ya_no_este_vigente(self):
        regla = ComisionColaboradorLinea.objects.get(colaborador=self.alberto)
        regla.vigente_hasta = date(2026, 9, 5)  # una cobertura que ya terminó
        regla.save()
        self._devolver(self.luis, self.antibiotico, 10)

        comisiones, _ = self._comisiones(fecha_desde="2026-09-10", fecha_hasta="2026-09-16")

        self.assertEqual(comisiones["alberto"], Decimal("-40.00"))

    def test_el_excel_trae_las_devoluciones_descontadas_y_las_posteriores(self):
        import io

        self._devolver(self.luis, self.antibiotico, 10)
        textos = {}
        for desde, hasta in [("2026-08-27", "2026-09-02"), ("2026-09-10", "2026-09-16")]:
            respuesta = self.client.get(
                reverse("comisiones_ruta:reporte-excel"), {"fecha_desde": desde, "fecha_hasta": hasta}
            )
            luis = openpyxl.load_workbook(io.BytesIO(respuesta.content))["LUIS"]
            textos[desde] = [str(c.value) for fila in luis.iter_rows() for c in fila if c.value is not None]

        self.assertIn("Se descuenta en el periodo", textos["2026-08-27"])
        self.assertIn("10/09/2026 al 16/09/2026", textos["2026-08-27"])
        self.assertIn("Devoluciones descontadas en este periodo", textos["2026-09-10"])
        self.assertIn("27/08/2026 al 02/09/2026", textos["2026-09-10"])
        self.assertIn("-50", textos["2026-09-10"])
