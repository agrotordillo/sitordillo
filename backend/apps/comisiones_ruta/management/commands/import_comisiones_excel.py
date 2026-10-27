"""Carga las reglas de comisión por colaborador (ComisionColaboradorLinea)
desde el Excel con el que se pagan las comisiones cada semana (por defecto
datos/COMISIONES P43.xlsx, el más reciente): una hoja por colaborador.

La mayoría de las hojas tienen un bloque por línea con sus productos, la base
(Precio 1 + 2 y Precio 3 + 4 + 5) y una fórmula `=+E<n>*X%` por cada base: la
primera es el `porcentaje` (PUBLICO/MEDIO MAYOREO) y la segunda el
`porcentaje_otras_listas`. Reglas, decididas con el usuario y validadas
contra las ventas de scv_web de cada semana (P36 y P43, 2026-10-27):
  - Al cargar un Excel se reemplazan las reglas que vinieron de un Excel
    para los colaboradores que trae; las capturadas a mano no se tocan.
  - La línea de cada bloque sale de sus productos, no del título: hay títulos
    copiados de otro bloque (p. ej. SEMILLAS Y FERTILIZANTES titulado
    REFACCIONES PARA BOMBAS). Se avisa cuando no coinciden. Los bloques que
    se repiten (la semana partida por el día de descanso, o un colaborador
    con dos usuarios) dan la misma regla.
  - Un bloque titulado con una categoría (SOBRES) es una regla de esa
    categoría; FARMACIA RESPONSIVA es la clase MEDICOS; ALIMENTOS en las
    hojas de los médicos son solo los "premium" (clases MASCOTAS*). Algunos
    bloques son sobre toda la venta de una sucursal (ALCANCE_SUCURSAL): así
    cuadran sus bases, no con lo que vendió el colaborador.
  - Las hojas de un solo total (HOJAS_DE_TOTAL) cobran un % de todo lo que
    vende el colaborador, incluido lo vendido con precio especial de sucursal
    (Precio 6-9), que en las demás no cuenta. MARLIT: solo ALIMENTOS en
    Precio 1-2.
  - JHONY cobra distinto según dónde vende (ESQUEMA_POR_SUCURSAL).
  - LUIS GERARDO y MIGUEL (o quien lo cubre, como TARACENA en sus
    vacaciones) cobran cada línea sobre toda la venta de LERDO de su clase de
    producto (ESQUEMAS_POR_CLASE), salvo su día de descanso (DESCANSOS). Las
    hojas de cobertura (COBERTURAS) solo valen en las fechas de su encabezado.
    MIGUEL no viene en el P43: su esquema se carga del P36 con
    --hojas "MIGUEL RMZ".
  - Los bloques de devoluciones no son reglas: el reporte ya descuenta lo
    devuelto. TARACENA, WILLIAMS y ZAMORA no venden en scv_web.
"""
import re
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.comisiones_ruta.models import TODOS_LOS_DIAS, ComisionColaboradorLinea
from apps.products.models import Almacen, Categoria, Clase, Linea, Producto

ARCHIVO = Path(settings.BASE_DIR).parent / "datos" / "COMISIONES P43.xlsx"
MARCA = "Cargado de "  # observaciones de las reglas que vienen de un Excel
# Hoja del Excel -> usuario de sitordillo (los vendedores de scv_web, ver
# apps.core.scv_web.VendedoresLegado). Incluye los nombres de hoja de P36 y P43.
HOJA_USUARIO = {
    "GILBERTO": "gilberto",
    "ANIBAL": "anibalp",
    "JUAN DANIEL": "jdaniel",
    "JULIO Y MENESES": "lmeneses",
    "JULIO Y L.GMENESES": "lmeneses",
    "MVZ ALBERTO": "alberto",
    "MVZ GUSTAVO": "mvzgustavo",
    "MIGUEL RMZ": "miguelr",
    "LUIS GDO": "luisg",
    "GUILLERMO": "guillermo",
    "JHONY": "jhony",
    "FRANKLIN": "frank",
    "MARLIT": "mgonzalez",
    "JUAN CARLOS": "jcarlos",
}
# Hojas con una sola fórmula sobre el total: hoja -> (línea o None = todas,
# % de las otras listas igual al principal o 0, incluye precio especial).
HOJAS_DE_TOTAL = {
    "JUAN DANIEL": (None, True, True),
    "FRANKLIN": (None, True, True),
    "JUAN CARLOS": (None, True, True),
    "MARLIT": ("ALIMENTOS", False, False),
}
# Hojas que cobran distinto según la sucursal. El Excel no dice de qué
# sucursal es cada bloque; se distinguen así (validado con P43): la fórmula
# sobre el total (antes del primer bloque) es "total"; los bloques al 1%,
# "uno_por_ciento"; los demás (5%, 2%, 10%), "otro_por_ciento". Cada grupo:
# sucursal y, si es sobre toda su venta, la clase y los días (JHONY cubrió el
# esquema de MIGUEL en LERDO el domingo de esa semana).
ESQUEMA_POR_SUCURSAL = {
    "JHONY": {
        "total": {"almacen": "IQUINUAPA"},
        "uno_por_ciento": {"almacen": "TECOLUTILLA"},
        "otro_por_ciento": {"almacen": "LERDO", "toda_la_sucursal": True, "clase": "MIGUEL RAMIREZ",
                            "dias": "6", "vigencia_de_la_hoja": True},
    },
}
# Hojas cuyas líneas se calculan sobre toda la venta de una sucursal de una
# clase de producto (el reporte de scv_web se "agrupa por clase": documento
# INFORMACION SCVWEB). TARACENA cubrió las vacaciones de MIGUEL.
ESQUEMAS_POR_CLASE = {
    "LUIS GDO": ("LERDO", "JOSE MANUEL"),
    "MIGUEL RMZ": ("LERDO", "MIGUEL RAMIREZ"),
    "JUAN TARACENA": ("LERDO", "MIGUEL RAMIREZ"),
}
# Día de descanso (0 = lunes ... 6 = domingo) de quien cobra sobre toda la
# venta de una sucursal: ese día no le toca (lo cobra quien lo cubre). Del
# documento INFORMACION SCVWEB; TARACENA no cubrió el domingo (lo cubrió JHONY).
DESCANSOS = {"MVZ ALBERTO": 2, "LUIS GDO": 3, "MIGUEL RMZ": 4, "JUAN TARACENA": 6}
# Hojas de quien cubre a otro: sus reglas solo valen en las fechas del
# encabezado de la hoja ("COMISIONES DEL 15 AL 21 DE OCTUBRE 2026").
COBERTURAS = {"JUAN TARACENA", "MVZ GUSTAVO"}
# Quien cobra comisión pero no vende con un usuario propio en scv_web: se le
# crea un usuario inactivo (como a los vendedores, ver scv_web.VendedoresLegado).
USUARIOS_NUEVOS = {"JUAN TARACENA": ("jtaracena", "JUAN ALBERTO TARACENA")}
# Títulos de bloque que no son línea ni categoría: qué clase de producto son.
GRUPOS_POR_CLASE = {"FARMACIA RESPONSIVA": ["MEDICOS"]}
# En estas hojas, ALIMENTOS son solo los alimentos "premium" para mascotas.
ALIMENTOS_PREMIUM = {"MVZ ALBERTO", "LUIS GDO"}
CLASES_PREMIUM = ["MASCOTAS ADM", "MASCOTAS NUPEC", "MASCOTAS FLAGASA", "MASCOTAS-VIMIFOS"]
# (hoja, título del bloque) -> sucursal de cuya venta total sale la base.
ALCANCE_SUCURSAL = {
    ("MVZ ALBERTO", "FARMACIA RESPONSIVA"): "LERDO",
    ("MVZ GUSTAVO", "FARMACIA RESPONSIVA"): "LERDO",
    ("LUIS GDO", "FARMACIA RESPONSIVA"): "LERDO",
    ("LUIS GDO", "SOBRES"): "LERDO",
}
IGNORAR_PREFIJO = "DEVOLUCION"

PORCENTAJE = re.compile(r"\*\s*([\d.]+)\s*%")
AJUSTE_A_MANO = re.compile(r"%\s*([+-]\s*[\d.]+)")
SUFIJO_SKU = re.compile(r"\s*\([^)]*\)\s*$")
UNO = Decimal("1")
MESES = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8, "SEP": 9, "OCT": 10,
         "NOV": 11, "DIC": 12}


def normalizar(valor):
    valor = re.sub(r"\bP/\s*", "PARA ", str(valor or "").upper())
    return re.sub(r"\s+", " ", valor).strip()


def periodo_de_la_hoja(hoja):
    """(desde, hasta) del encabezado "COMISIONES DEL 15 AL 21 DE OCTUBRE 2026
    P43" (también "DEL 27 DE AGOSTO AL 02 DE SEPTIEMBRE 2026" o "DEL 18 DE
    OCTUBRE 2026"), o None si no se entiende."""
    for fila in hoja.iter_rows(min_row=1, max_row=3, values_only=True):
        for valor in fila:
            texto = normalizar(valor)
            if "COMISIONES DEL" not in texto:
                continue
            texto = texto.split("COMISIONES DEL", 1)[1]
            anio = re.findall(r"\b(20\d{2})\b", texto)
            if not anio:
                return None
            texto = re.sub(r"\bP\d+\b|\b20\d{2}\b", " ", texto)
            pares = re.findall(r"\b(\d{1,2})\b(?:\s+DE\s+([A-ZÑ]+))?", texto)
            if not pares:
                return None
            meses, siguiente = [], None
            for _, mes in reversed(pares):  # un día sin mes toma el del siguiente
                siguiente = MESES.get(mes[:3], siguiente) if mes else siguiente
                meses.append(siguiente)
            meses.reverse()
            if None in (meses[0], meses[-1]):
                return None
            try:
                return (
                    date(int(anio[-1]), meses[0], int(pares[0][0])),
                    date(int(anio[-1]), meses[-1], int(pares[-1][0])),
                )
            except ValueError:
                return None
    return None


def formulas_de_porcentaje(celdas):
    return sorted(
        (c for c in celdas if isinstance(c.value, str) and c.value.startswith("=") and "%" in c.value),
        key=lambda c: c.column,
    )


def porcentajes(formulas):
    return [Decimal(PORCENTAJE.search(c.value).group(1)) for c in formulas if PORCENTAJE.search(c.value)]


def leer_hoja(hoja):
    """(fórmulas antes del primer bloque, [{titulo, productos, porcentajes}]).
    Un bloque empieza en un renglón con un solo texto seguido del encabezado
    "Registros"; los productos van en la columna anterior a "Registros"."""
    filas = list(hoja.iter_rows())
    previas, bloques, actual, columna_nombre = [], [], None, None
    for i, fila in enumerate(filas):
        llenas = [c for c in fila if c.value is not None]
        siguiente = filas[i + 1] if i + 1 < len(filas) else ()
        registros = [c.column for c in siguiente if str(c.value or "").strip() == "Registros"]
        if len(llenas) == 1 and isinstance(llenas[0].value, str) and registros:
            titulo = llenas[0].value.strip()
            if normalizar(titulo).startswith("COMISIONES"):
                # El resumen del encabezado ("COMISIONES DEL ... P43"): su
                # fórmula es la del total, no un bloque de línea.
                actual = None
                continue
            actual = {"titulo": titulo, "productos": [], "porcentajes": []}
            columna_nombre = registros[0] - 1
            bloques.append(actual)
            continue
        formulas = formulas_de_porcentaje(llenas)
        if actual is None:
            previas.extend(formulas)
            continue
        if formulas:
            actual["porcentajes"] = porcentajes(formulas)
            continue
        nombre = fila[columna_nombre - 1].value if columna_nombre and len(fila) >= columna_nombre else None
        if isinstance(nombre, str) and nombre.strip() not in ("Registros", "Total"):
            actual["productos"].append(nombre.strip())
    return previas, [b for b in bloques if b["porcentajes"]]


def dias_sin(descanso):
    """Los días de la semana (texto "0123456") sin el de descanso."""
    return TODOS_LOS_DIAS if descanso is None else TODOS_LOS_DIAS.replace(str(descanso), "")


class Command(BaseCommand):
    help = "Carga las reglas de comisión por colaborador desde el Excel de comisiones (por defecto el P43)."

    def add_arguments(self, parser):
        parser.add_argument("--archivo", default=str(ARCHIVO))
        parser.add_argument(
            "--hojas", nargs="+", default=None,
            help='Carga solo estas hojas (p. ej. --hojas "MIGUEL RMZ" del P36, que no viene en el P43).',
        )
        parser.add_argument("--dry-run", action="store_true", help="No guarda nada, solo muestra la tabla.")

    def handle(self, *args, **options):
        ruta = Path(options["archivo"])
        if not ruta.exists():
            raise CommandError(f"No encuentro {ruta}")
        self.dry_run = options["dry_run"]
        # Abierto con `with`: openpyxl no suelta el archivo por sí solo y en
        # Windows lo dejaría bloqueado (p. ej. para Excel).
        with open(ruta, "rb") as archivo:
            libro = openpyxl.load_workbook(archivo, data_only=False)
        hojas = libro.sheetnames
        if options["hojas"]:
            faltan = [h for h in options["hojas"] if h not in hojas]
            if faltan:
                raise CommandError(f"El Excel no trae estas hojas: {', '.join(faltan)}")
            hojas = options["hojas"]

        self.lineas = {normalizar(l.nombre): l for l in Linea.objects.all()}
        self.categorias = {normalizar(c.nombre): c for c in Categoria.objects.all()}
        self.clases = {normalizar(c.nombre): c for c in Clase.objects.all()}
        self.almacenes = {normalizar(a.nombre): a for a in Almacen.objects.all()}
        self.linea_de_producto = {
            normalizar(SUFIJO_SKU.sub("", nombre)): linea_id
            for nombre, linea_id in Producto.objects.exclude(linea=None).values_list("nombre", "linea_id")
        }
        self.lineas_por_id = {l.pk: l for l in self.lineas.values()}
        nombres_usuario = {**HOJA_USUARIO, **{h: u for h, (u, _) in USUARIOS_NUEVOS.items()}}
        usuarios = {u.username: u for u in get_user_model().objects.filter(username__in=nombres_usuario.values())}

        self.reglas, self.avisos, cargados = {}, [], set()
        for nombre_hoja in hojas:
            usuario = usuarios.get(nombres_usuario.get(nombre_hoja)) or self._usuario_nuevo(nombre_hoja)
            if usuario is None:
                self.avisos.append(f"{nombre_hoja}: no vende en scv_web (sin usuario de vendedor), se omite")
                continue
            if usuario.pk:
                cargados.add(usuario.pk)
            hoja = libro[nombre_hoja]
            self.periodo = None
            if nombre_hoja in COBERTURAS or nombre_hoja in ESQUEMA_POR_SUCURSAL:
                self.periodo = periodo_de_la_hoja(hoja)
                if self.periodo is None and nombre_hoja in COBERTURAS:
                    self.avisos.append(f"{nombre_hoja}: no se entiende el periodo del encabezado, se omite")
                    continue
            previas, bloques = leer_hoja(hoja)
            if nombre_hoja in HOJAS_DE_TOTAL:
                self._hoja_de_total(nombre_hoja, previas, bloques, usuario)
            elif nombre_hoja in ESQUEMA_POR_SUCURSAL:
                self._hoja_por_sucursal(nombre_hoja, previas, bloques, usuario)
            else:
                for bloque in bloques:
                    self._bloque(nombre_hoja, bloque, usuario)

        for regla in sorted(self.reglas.values(), key=lambda r: (r.colaborador.username, r.descripcion)):
            self.stdout.write(
                f"  {regla.colaborador.username:12} {regla.descripcion:90} {regla.porcentaje}% / {regla.porcentaje_otras_listas}%"
            )
        for aviso in self.avisos:
            self.stdout.write(self.style.WARNING(f"  - {aviso}"))

        anteriores = ComisionColaboradorLinea.objects.filter(
            colaborador_id__in=cargados, observaciones__startswith=MARCA
        )
        if self.dry_run:
            self.stdout.write(self.style.WARNING(
                f"Dry-run: se reemplazarían {anteriores.count()} reglas cargadas de un Excel por {len(self.reglas)}."
            ))
            return
        with transaction.atomic():
            borradas = anteriores.delete()[0]
            for regla in self.reglas.values():
                regla.observaciones = f"{MARCA}{ruta.name}"
                try:
                    regla.full_clean()
                except ValidationError as exc:
                    raise CommandError(
                        f"{regla.colaborador.username} · {regla.descripcion}: {'; '.join(exc.messages)} "
                        "(choca con una regla capturada a mano). No se cargó nada."
                    ) from exc
                regla.save()
        self.stdout.write(self.style.SUCCESS(
            f"Listo: {len(self.reglas)} reglas cargadas de {ruta.name} (reemplazan {borradas} de un Excel anterior)."
        ))

    def _usuario_nuevo(self, nombre_hoja):
        """El usuario inactivo de quien cobra comisión sin vender con usuario
        propio (USUARIOS_NUEVOS); en el dry-run no se guarda."""
        if nombre_hoja not in USUARIOS_NUEVOS:
            return None
        username, nombre = USUARIOS_NUEVOS[nombre_hoja]
        usuario = get_user_model()(username=username, first_name=nombre, is_active=False)
        usuario.set_unusable_password()
        if not self.dry_run:
            usuario.save()
        self.avisos.append(f"{nombre_hoja}: se {'crearía' if self.dry_run else 'creó'} el usuario inactivo {username}")
        return usuario

    def _hoja_de_total(self, nombre_hoja, previas, bloques, usuario):
        linea_nombre, otras_igual, especial = HOJAS_DE_TOTAL[nombre_hoja]
        pcts = porcentajes(previas) + [p for b in bloques for p in b["porcentajes"]]
        if len(pcts) != 1:
            self.avisos.append(f"{nombre_hoja}: se esperaba una sola fórmula de % y hay {len(pcts)}, se omite")
            return
        [pct] = pcts
        linea = self.lineas[normalizar(linea_nombre)] if linea_nombre else None
        self._agregar(nombre_hoja, usuario, pct, pct if otras_igual else Decimal("0"), linea=linea, especial=especial)

    def _hoja_por_sucursal(self, nombre_hoja, previas, bloques, usuario):
        esquema = ESQUEMA_POR_SUCURSAL[nombre_hoja]
        total = porcentajes(previas)
        for celda in previas:
            ajuste = AJUSTE_A_MANO.search(celda.value)
            if ajuste:
                self.avisos.append(
                    f"{nombre_hoja}: la fórmula del total suma a mano {ajuste.group(1).replace(' ', '')} "
                    f"({celda.value}); eso no se puede cargar como regla"
                )
        if total:
            self._agregar(nombre_hoja, usuario, total[0], total[0], especial=True, **self._grupo(esquema["total"]))
        for bloque in bloques:
            clave = "uno_por_ciento" if bloque["porcentajes"][0] == UNO else "otro_por_ciento"
            self._bloque(nombre_hoja, bloque, usuario, **self._grupo(esquema[clave]))

    def _grupo(self, grupo):
        """Alcance, sucursal, clase, días y vigencia de un grupo de
        ESQUEMA_POR_SUCURSAL."""
        Alcance = ComisionColaboradorLinea.Alcance
        datos = {
            "alcance": Alcance.SUCURSAL if grupo.get("toda_la_sucursal") else Alcance.PROPIAS_EN_SUCURSAL,
            "almacen": self.almacenes[normalizar(grupo["almacen"])],
            "dias": grupo.get("dias"),
        }
        if grupo.get("clase"):
            datos["clase_del_esquema"] = self.clases.get(normalizar(grupo["clase"]))
        if grupo.get("vigencia_de_la_hoja") and self.periodo:
            datos["vigencia"] = self.periodo
        return datos

    def _bloque(self, nombre_hoja, bloque, usuario, alcance=None, almacen=None, dias=None, vigencia=None,
                clase_del_esquema=None):
        titulo = normalizar(bloque["titulo"])
        if titulo.startswith(IGNORAR_PREFIJO):
            return
        variable, *resto = bloque["porcentajes"]
        otras = resto[0] if resto else variable
        if (nombre_hoja, titulo) in ALCANCE_SUCURSAL:
            alcance = ComisionColaboradorLinea.Alcance.SUCURSAL
            almacen = self.almacenes[normalizar(ALCANCE_SUCURSAL[(nombre_hoja, titulo)])]
        extra = {"alcance": alcance, "almacen": almacen, "dias": dias, "vigencia": vigencia}

        if titulo in GRUPOS_POR_CLASE:
            for clase in GRUPOS_POR_CLASE[titulo]:
                self._agregar(nombre_hoja, usuario, variable, otras, clase=self.clases.get(normalizar(clase)), **extra)
            return
        if titulo not in self.lineas and titulo in self.categorias:
            self._agregar(nombre_hoja, usuario, variable, otras, categoria=self.categorias[titulo], **extra)
            return
        linea_titulo = self.lineas.get(titulo)
        if linea_titulo is None:
            self.avisos.append(f"{nombre_hoja}: '{bloque['titulo']}' no es línea, categoría ni grupo conocido, se omite")
            return
        votos = Counter(
            self.linea_de_producto[normalizar(p)] for p in bloque["productos"] if normalizar(p) in self.linea_de_producto
        )
        linea = self.lineas_por_id[votos.most_common(1)[0][0]] if votos else linea_titulo
        if linea != linea_titulo:
            self.avisos.append(f"{nombre_hoja}: el bloque titulado '{bloque['titulo']}' trae productos de {linea.nombre}")
        if nombre_hoja in ALIMENTOS_PREMIUM and normalizar(linea.nombre) == "ALIMENTOS":
            for clase in CLASES_PREMIUM:
                if normalizar(clase) in self.clases:
                    self._agregar(nombre_hoja, usuario, variable, otras, linea=linea,
                                  clase=self.clases[normalizar(clase)], **extra)
            return
        if nombre_hoja in ESQUEMAS_POR_CLASE:
            # Línea de su clase vendida por cualquiera en la sucursal.
            sucursal, clase = ESQUEMAS_POR_CLASE[nombre_hoja]
            extra.update(alcance=ComisionColaboradorLinea.Alcance.SUCURSAL, almacen=self.almacenes[normalizar(sucursal)])
            clase_del_esquema = self.clases.get(normalizar(clase))
        self._agregar(nombre_hoja, usuario, variable, otras, linea=linea, clase=clase_del_esquema, **extra)

    def _agregar(self, nombre_hoja, usuario, porcentaje, otras, *, linea=None, categoria=None, clase=None,
                 almacen=None, alcance=None, especial=False, dias=None, vigencia=None):
        Alcance = ComisionColaboradorLinea.Alcance
        alcance = alcance or (Alcance.SUCURSAL if almacen else Alcance.PROPIAS)
        if dias is None:
            # Sobre toda la venta de una sucursal, el día de descanso lo cobra quien cubre.
            dias = dias_sin(DESCANSOS.get(nombre_hoja)) if alcance == Alcance.SUCURSAL else TODOS_LOS_DIAS
        if vigencia is None and nombre_hoja in COBERTURAS:
            vigencia = self.periodo
        desde, hasta = vigencia or (None, None)
        regla = ComisionColaboradorLinea(
            colaborador=usuario, alcance=alcance, almacen=almacen, linea=linea, categoria=categoria, clase=clase,
            porcentaje=porcentaje, porcentaje_otras_listas=otras, comisiona_precio_especial=especial,
            dias_semana=dias, vigente_desde=desde, vigente_hasta=hasta,
        )
        clave = (usuario.username, alcance, getattr(almacen, "pk", None), getattr(linea, "pk", None),
                 getattr(categoria, "pk", None), getattr(clase, "pk", None), desde)
        anterior = self.reglas.get(clave)
        if anterior is not None and (anterior.porcentaje, anterior.porcentaje_otras_listas) != (porcentaje, otras):
            self.avisos.append(
                f"{nombre_hoja} · {regla.descripcion}: dos bloques con % distinto "
                f"({anterior.porcentaje}% y {porcentaje}%), se queda el primero"
            )
            return
        self.reglas.setdefault(clave, regla)

