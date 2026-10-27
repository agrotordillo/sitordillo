"""Carga las reglas de comisión por colaborador (ComisionColaboradorLinea)
desde el Excel con el que se pagan hoy las comisiones (datos/COMISIONES
P36.xlsx): una hoja por colaborador.

La mayoría de las hojas tienen un bloque por línea con sus productos, la base
(Precio 1 + 2 y Precio 3 + 4 + 5) y una fórmula `=+E<n>*X%` por cada base: la
primera es el `porcentaje` (PUBLICO/MEDIO MAYOREO) y la segunda el
`porcentaje_otras_listas`. Reglas, decididas con el usuario y validadas
contra las ventas de scv_web de esa semana (2026-10-27):
  - Se carga tal cual. La línea de cada bloque sale de sus productos, no del
    título: hay títulos copiados de otro bloque (p. ej. en GUILLERMO un bloque
    de SEMILLAS Y FERTILIZANTES titulado REFACCIONES PARA BOMBAS).
  - Un bloque titulado con una categoría (SOBRES) es una regla de esa
    categoría; FARMACIA RESPONSIVA es la clase MEDICOS. En la hoja del MVZ
    esa regla es sobre toda la venta de LERDO (ver ALCANCE_SUCURSAL): así
    cuadra su base, no con lo que él vendió.
  - JHONY y FRANKLIN cobran un % de todo lo que venden, en cualquier lista
    (hoja de resumen con una sola fórmula "TOTAL A PERCIBIR").
  - MARLIT cobra solo ALIMENTOS en Precio 1-2 (su hoja es una lista de notas
    con una sola fórmula): 0% en las otras listas.
  - DEVOLUCION AGROQUIMICOS no es una regla: el reporte ya descuenta lo
    devuelto. TARACENA, WILLIAMS y ZAMORA no venden en scv_web.
"""
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.comisiones_ruta.models import ComisionColaboradorLinea
from apps.products.models import Almacen, Categoria, Clase, Linea, Producto

ARCHIVO = Path(settings.BASE_DIR).parent / "datos" / "COMISIONES P36.xlsx"
# Hoja del Excel -> usuario de sitordillo (los vendedores de scv_web, ver
# apps.core.scv_web.VendedoresLegado).
HOJA_USUARIO = {
    "GILBERTO": "gilberto",
    "ANIBAL": "anibalp",
    "JUAN DANIEL": "jdaniel",
    "JULIO Y MENESES": "lmeneses",
    "MVZ ALBERTO": "alberto",
    "MIGUEL RMZ": "miguelr",
    "LUIS GDO": "luisg",
    "GUILLERMO": "guillermo",
    "JHONY": "jhony",
    "FRANKLIN": "frank",
    "MARLIT": "mgonzalez",
}
# Hojas con una sola fórmula sobre el total: hoja -> línea (None = todas) y
# si las otras listas comisionan igual (True) o nada (False).
HOJAS_DE_TOTAL = {
    "JHONY": (None, True),
    "FRANKLIN": (None, True),
    "MARLIT": ("ALIMENTOS", False),
}
# Títulos de bloque que no son línea ni categoría: qué clase de producto son.
GRUPOS_POR_CLASE = {"FARMACIA RESPONSIVA": "MEDICOS"}
# (hoja, título del bloque) -> sucursal de cuya venta total sale la base.
ALCANCE_SUCURSAL = {("MVZ ALBERTO", "FARMACIA RESPONSIVA"): "LERDO"}
IGNORAR = {"DEVOLUCION AGROQUIMICOS"}

PORCENTAJE = re.compile(r"\*\s*([\d.]+)\s*%")
SUFIJO_SKU = re.compile(r"\s*\([^)]*\)\s*$")


def normalizar(valor):
    valor = re.sub(r"\bP/\s*", "PARA ", str(valor or "").upper())
    return re.sub(r"\s+", " ", valor).strip()


def formulas_de_porcentaje(celdas):
    formulas = sorted(
        (c for c in celdas if isinstance(c.value, str) and c.value.startswith("=") and "%" in c.value),
        key=lambda c: c.column,
    )
    return [Decimal(PORCENTAJE.search(c.value).group(1)) for c in formulas if PORCENTAJE.search(c.value)]


def leer_bloques(hoja):
    """[{titulo, productos, porcentajes}] de una hoja. El bloque empieza en un
    renglón con un solo texto seguido del encabezado "Registros"; los
    productos van en la columna anterior a "Registros"."""
    filas = list(hoja.iter_rows())
    bloques, actual, columna_nombre = [], None, None
    for i, fila in enumerate(filas):
        llenas = [c for c in fila if c.value is not None]
        siguiente = filas[i + 1] if i + 1 < len(filas) else ()
        registros = [c.column for c in siguiente if str(c.value or "").strip() == "Registros"]
        if len(llenas) == 1 and isinstance(llenas[0].value, str) and registros:
            actual = {"titulo": llenas[0].value.strip(), "productos": [], "porcentajes": []}
            columna_nombre = registros[0] - 1
            bloques.append(actual)
            continue
        if actual is None:
            continue
        porcentajes = formulas_de_porcentaje(llenas)
        if porcentajes:
            actual["porcentajes"] = porcentajes
            continue
        nombre = fila[columna_nombre - 1].value if columna_nombre and len(fila) >= columna_nombre else None
        if isinstance(nombre, str) and nombre.strip() not in ("Registros", "Total"):
            actual["productos"].append(nombre.strip())
    return [b for b in bloques if b["porcentajes"]]


class Command(BaseCommand):
    help = "Carga las reglas de comisión por colaborador desde datos/COMISIONES P36.xlsx."

    def add_arguments(self, parser):
        parser.add_argument("--archivo", default=str(ARCHIVO))
        parser.add_argument("--dry-run", action="store_true", help="No guarda nada, solo muestra la tabla.")

    def handle(self, *args, **options):
        ruta = Path(options["archivo"])
        if not ruta.exists():
            raise CommandError(f"No encuentro {ruta}")
        # Abierto con `with`: openpyxl no suelta el archivo por sí solo y en
        # Windows lo dejaría bloqueado (p. ej. para Excel).
        with open(ruta, "rb") as archivo:
            libro = openpyxl.load_workbook(archivo, data_only=False)

        self.lineas = {normalizar(l.nombre): l for l in Linea.objects.all()}
        self.categorias = {normalizar(c.nombre): c for c in Categoria.objects.all()}
        self.clases = {normalizar(c.nombre): c for c in Clase.objects.all()}
        self.almacenes = {normalizar(a.nombre): a for a in Almacen.objects.all()}
        self.linea_de_producto = {
            normalizar(SUFIJO_SKU.sub("", nombre)): linea_id
            for nombre, linea_id in Producto.objects.exclude(linea=None).values_list("nombre", "linea_id")
        }
        self.lineas_por_id = {l.pk: l for l in self.lineas.values()}
        usuarios = {u.username: u for u in get_user_model().objects.filter(username__in=HOJA_USUARIO.values())}

        self.reglas, self.avisos = {}, []
        for nombre_hoja in libro.sheetnames:
            usuario = usuarios.get(HOJA_USUARIO.get(nombre_hoja))
            if usuario is None:
                self.avisos.append(f"{nombre_hoja}: no vende en scv_web (sin usuario de vendedor), se omite")
                continue
            hoja = libro[nombre_hoja]
            if nombre_hoja in HOJAS_DE_TOTAL:
                self._hoja_de_total(nombre_hoja, hoja, usuario)
            else:
                for bloque in leer_bloques(hoja):
                    self._bloque(nombre_hoja, bloque, usuario)

        for regla in sorted(self.reglas.values(), key=lambda r: (r.colaborador.username, r.descripcion)):
            self.stdout.write(
                f"  {regla.colaborador.username:12} {regla.descripcion:55} {regla.porcentaje}% / {regla.porcentaje_otras_listas}%"
            )
        for aviso in self.avisos:
            self.stdout.write(self.style.WARNING(f"  - {aviso}"))

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING(f"Dry-run: se cargarían {len(self.reglas)} reglas."))
            return
        with transaction.atomic():
            for regla in self.reglas.values():
                ComisionColaboradorLinea.objects.update_or_create(
                    colaborador=regla.colaborador, almacen=regla.almacen, linea=regla.linea,
                    categoria=regla.categoria, clase=regla.clase,
                    defaults={
                        "porcentaje": regla.porcentaje,
                        "porcentaje_otras_listas": regla.porcentaje_otras_listas,
                        "observaciones": "Cargado de COMISIONES P36.xlsx",
                    },
                )
        self.stdout.write(self.style.SUCCESS(f"Listo: {len(self.reglas)} reglas cargadas."))

    def _hoja_de_total(self, nombre_hoja, hoja, usuario):
        linea_nombre, otras_igual = HOJAS_DE_TOTAL[nombre_hoja]
        porcentajes = [p for fila in hoja.iter_rows() for p in formulas_de_porcentaje(fila)]
        if len(porcentajes) != 1:
            self.avisos.append(f"{nombre_hoja}: se esperaba una sola fórmula de % y hay {len(porcentajes)}, se omite")
            return
        [pct] = porcentajes
        linea = self.lineas[normalizar(linea_nombre)] if linea_nombre else None
        self._agregar(nombre_hoja, usuario, pct, pct if otras_igual else Decimal("0"), linea=linea)

    def _bloque(self, nombre_hoja, bloque, usuario):
        titulo = normalizar(bloque["titulo"])
        if titulo in IGNORAR:
            return
        variable, *resto = bloque["porcentajes"]
        otras = resto[0] if resto else variable
        almacen = None
        if (nombre_hoja, titulo) in ALCANCE_SUCURSAL:
            almacen = self.almacenes[normalizar(ALCANCE_SUCURSAL[(nombre_hoja, titulo)])]

        if titulo in GRUPOS_POR_CLASE:
            clase = self.clases.get(normalizar(GRUPOS_POR_CLASE[titulo]))
            self._agregar(nombre_hoja, usuario, variable, otras, clase=clase, almacen=almacen)
            return
        if titulo not in self.lineas and titulo in self.categorias:
            self._agregar(nombre_hoja, usuario, variable, otras, categoria=self.categorias[titulo], almacen=almacen)
            return
        linea_titulo = self.lineas.get(titulo)
        if linea_titulo is None:
            self.avisos.append(f"{nombre_hoja}: '{bloque['titulo']}' no es línea, categoría ni grupo conocido, se omite")
            return
        votos = Counter(
            self.linea_de_producto[normalizar(p)] for p in bloque["productos"] if normalizar(p) in self.linea_de_producto
        )
        linea = self.lineas_por_id[votos.most_common(1)[0][0]] if votos else linea_titulo
        if linea != linea_titulo:
            self.avisos.append(f"{nombre_hoja}: el bloque titulado '{bloque['titulo']}' trae productos de {linea.nombre}")
        self._agregar(nombre_hoja, usuario, variable, otras, linea=linea, almacen=almacen)

    def _agregar(self, nombre_hoja, usuario, porcentaje, otras, *, linea=None, categoria=None, clase=None, almacen=None):
        regla = ComisionColaboradorLinea(
            colaborador=usuario, almacen=almacen, linea=linea, categoria=categoria, clase=clase,
            porcentaje=porcentaje, porcentaje_otras_listas=otras,
        )
        clave = (usuario.pk, getattr(almacen, "pk", None), getattr(linea, "pk", None),
                 getattr(categoria, "pk", None), getattr(clase, "pk", None))
        anterior = self.reglas.get(clave)
        if anterior is not None and (anterior.porcentaje, anterior.porcentaje_otras_listas) != (porcentaje, otras):
            self.avisos.append(
                f"{nombre_hoja} · {regla.descripcion}: dos bloques con % distinto "
                f"({anterior.porcentaje}% y {porcentaje}%), se queda el primero"
            )
            return
        self.reglas.setdefault(clave, regla)
