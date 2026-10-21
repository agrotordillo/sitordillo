"""Actualización masiva de costo y precios con una hoja de cálculo, como
"Actualizar precios" del sistema anterior (control_ven_actualizar_precios):

- La plantilla es la misma: en A1 el título "Actualización de precios", en
  la fila 2 los encabezados codigo_producto, nombre_producto, costo,
  precio_1..precio_N, y desde la fila 3 un producto por renglón. Cada
  precio_N es el precio CON impuesto de la lista en la posición N del orden
  canónico (PRECIO_POSICIONES); una hoja "Instrucciones" dice cuál es cuál.
- Una celda vacía o en 0 no cambia nada. El costo vacío deja el que tiene
  el producto.
- El % de utilidad de cada precio que viene en el archivo se recalcula con
  el costo (el nuevo si viene); los precios que no vienen no se tocan,
  aunque cambie el costo (igual que el sistema anterior).
- A diferencia del anterior, primero se revisa todo el archivo y se muestra
  qué cambiaría y los errores por línea; solo se aplica si no hay errores y
  el usuario lo confirma (ver views/precios_masivos_views.py)."""

from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl
from django.db import transaction
from django.db.models.functions import Upper
from openpyxl.styles import Font

from .columnas_listado import _filtro_posicion, _titulo_posicion
from .models import ALMACEN_CEDIS, PRECIO_POSICIONES, Almacen, ListaPrecio, Producto, ProductoPrecio

TITULO = "Actualización de precios"
COLUMNA_CODIGO = "codigo_producto"
COLUMNA_NOMBRE = "nombre_producto"
COLUMNA_COSTO = "costo"
COLUMNAS_PRECIO = {f"precio_{posicion}": (posicion, lista, almacen) for posicion, lista, almacen in PRECIO_POSICIONES}
ENCABEZADOS = [COLUMNA_CODIGO, COLUMNA_NOMBRE, COLUMNA_COSTO, *COLUMNAS_PRECIO]
MAX_RENGLONES = 20_000
UTILIDAD_MAXIMA = Decimal("999.9999")


def generar_plantilla():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Precios"
    hoja["A1"] = TITULO
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(ENCABEZADOS)
    for celda in hoja[2]:
        celda.font = Font(bold=True)
    for indice, encabezado in enumerate(ENCABEZADOS, start=1):
        letra = hoja.cell(row=2, column=indice).column_letter
        hoja.column_dimensions[letra].width = 40 if encabezado == COLUMNA_NOMBRE else 16
    # El código como texto: así Excel no le quita los ceros a la izquierda
    # ("03") ni lo convierte a número.
    hoja.column_dimensions["A"].number_format = "@"

    instrucciones = libro.create_sheet("Instrucciones")
    renglones = [
        ["Cómo llenar la hoja Precios"],
        [],
        ["codigo_producto", "SKU (código) del producto. Obligatorio."],
        ["nombre_producto", "Solo como referencia; no se actualiza."],
        ["costo", "Precio de costo nuevo. Vacío o 0 = no cambia."],
        *[[columna, f"Precio con impuesto de {_titulo_posicion(lista, almacen)}. Vacío o 0 = no cambia."]
          for columna, (_, lista, almacen) in COLUMNAS_PRECIO.items()],
        [],
        ["El % de utilidad de cada precio capturado se recalcula con el costo."],
        ["Los precios que dejes vacíos no cambian, aunque cambie el costo."],
        ["No borres ni cambies las filas 1 y 2 de la hoja Precios."],
    ]
    for renglon in renglones:
        instrucciones.append(renglon)
    instrucciones["A1"].font = Font(bold=True, size=13)
    instrucciones.column_dimensions["A"].width = 20
    instrucciones.column_dimensions["B"].width = 80
    return libro


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _numero(valor, decimales):
    """Decimal redondeado, None si viene vacío o en 0 (= no cambia).
    ValueError si no es un número o es negativo."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        if isinstance(valor, (int, float)):
            numero = Decimal(str(valor))
        else:
            numero = Decimal(str(valor).replace("$", "").replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f'"{valor}" no es un número')
    if not numero.is_finite():
        raise ValueError(f'"{valor}" no es un número')
    if numero < 0:
        raise ValueError("no puede ser negativo")
    if numero == 0:
        return None
    return numero.quantize(Decimal(1).scaleb(-decimales), rounding=ROUND_HALF_UP)


def utilidad_de(producto, costo, precio_con_impuesto):
    """% de utilidad de un precio con impuesto sobre el costo, con las
    tasas del producto. None si no hay costo o se sale del rango del campo."""
    if not costo:
        return None
    tasa_iva = producto.tasa_iva / 100 if producto.tipo_iva == Producto.TipoIVA.GRAVADO else Decimal("0")
    tasa_ieps = producto.tasa_ieps / 100 if producto.aplica_ieps and producto.tasa_ieps else Decimal("0")
    sin_impuesto = precio_con_impuesto / ((1 + tasa_ieps) * (1 + tasa_iva))
    utilidad = ((sin_impuesto - costo) / costo * 100).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    return utilidad if abs(utilidad) <= UTILIDAD_MAXIMA else None


def _destinos_de_posiciones():
    """{posicion: (ListaPrecio, Almacen|None)} de las posiciones cuya lista
    (y sucursal) existen en el sistema."""
    listas = {lista.nombre: lista for lista in ListaPrecio.objects.all()}
    destinos = {}
    for posicion, nombre_lista, nombre_almacen in PRECIO_POSICIONES:
        lista = listas.get(nombre_lista)
        if lista is None:
            continue
        if nombre_almacen is None:
            destinos[posicion] = (lista, None)
            continue
        if nombre_almacen == ALMACEN_CEDIS:
            almacenes = list(Almacen.objects.filter(tipo=Almacen.Tipo.CEDIS)[:2])
        else:
            almacenes = list(Almacen.objects.filter(nombre=nombre_almacen)[:2])
        if len(almacenes) == 1:
            destinos[posicion] = (lista, almacenes[0])
    return destinos


def leer_archivo(archivo, *, puede_cambiar_costo):
    """Revisa el archivo sin guardar nada. Regresa (cambios, errores):
    `cambios` es una lista serializable (va a la sesión) con lo que cambiaría
    en cada producto; `errores` son textos "Línea N: ...". Un renglón sin
    ningún cambio respecto a lo que ya tiene el producto no se lista."""
    try:
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        return [], ["El archivo no es una hoja de cálculo .xlsx válida."]
    hoja = libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)

    primera = next(filas, None) or ()
    if not _texto(primera[0] if primera else None).startswith(TITULO):
        return [], [f'El archivo no tiene la estructura de la plantilla (en A1 debe decir "{TITULO}").']
    encabezados = [_texto(valor).lower() for valor in (next(filas, None) or ())]
    if COLUMNA_CODIGO not in encabezados:
        return [], [f'No se encontró la columna "{COLUMNA_CODIGO}" en la fila 2.']
    indice = {encabezado: posicion for posicion, encabezado in enumerate(encabezados) if encabezado}
    columnas_precio = [columna for columna in COLUMNAS_PRECIO if columna in indice]
    sobrantes = [e for e in indice if e.startswith("precio_") and e not in COLUMNAS_PRECIO]

    destinos = _destinos_de_posiciones()
    renglones = []
    errores = []
    for numero_linea, fila in enumerate(filas, start=3):
        if numero_linea - 2 > MAX_RENGLONES:
            errores.append((numero_linea, f"El archivo tiene más de {MAX_RENGLONES:,} renglones; divídelo en varios."))
            break
        celda = lambda columna: fila[indice[columna]] if columna in indice and indice[columna] < len(fila) else None
        codigo = _texto(celda(COLUMNA_CODIGO))
        valores = [celda(c) for c in (COLUMNA_COSTO, *columnas_precio, *sobrantes)]
        if not codigo and all(v in (None, "") for v in valores):
            continue
        if not codigo:
            errores.append((numero_linea, f"Línea {numero_linea}: falta el código de producto."))
            continue
        try:
            costo = _numero(celda(COLUMNA_COSTO), 4)
        except ValueError as error:
            errores.append((numero_linea, f"Línea {numero_linea} ({codigo}): el costo {error}."))
            continue
        precios = {}
        error_linea = None
        for columna in columnas_precio:
            try:
                valor = _numero(celda(columna), 2)
            except ValueError as error:
                error_linea = f"{columna} {error}"
                break
            if valor is not None:
                posicion = COLUMNAS_PRECIO[columna][0]
                if posicion not in destinos:
                    _, lista, almacen = COLUMNAS_PRECIO[columna]
                    error_linea = f"{columna}: no existe la lista {_titulo_posicion(lista, almacen)} en el sistema"
                    break
                precios[posicion] = valor
        for columna in sobrantes:
            if error_linea is None and _texto(celda(columna)) not in ("", "0"):
                error_linea = f"{columna} no corresponde a ninguna lista (solo hay precio_1 a precio_{len(COLUMNAS_PRECIO)})"
        if error_linea:
            errores.append((numero_linea, f"Línea {numero_linea} ({codigo}): {error_linea}."))
            continue
        renglones.append((numero_linea, codigo, costo, precios))

    # El código se compara sin distinguir mayúsculas/minúsculas.
    codigos = {codigo.upper() for _, codigo, _, _ in renglones}
    productos = {
        p.sku_mayusculas: p
        for p in Producto.objects.annotate(sku_mayusculas=Upper("sku")).filter(sku_mayusculas__in=codigos)
    }
    precios_actuales = {}
    for precio in ProductoPrecio.objects.filter(producto__in=productos.values()):
        precios_actuales[(precio.producto_id, precio.lista_precio_id, precio.almacen_id)] = precio.precio_con_impuesto

    cambios = []
    vistos = {}
    for numero_linea, codigo, costo, precios in renglones:
        producto = productos.get(codigo.upper())
        if producto is None:
            errores.append((numero_linea, f"Línea {numero_linea}: el código de producto {codigo} no existe."))
            continue
        if producto.pk in vistos:
            errores.append((numero_linea, f"Línea {numero_linea}: el código {codigo} ya viene en la línea {vistos[producto.pk]}."))
            continue
        vistos[producto.pk] = numero_linea
        if costo is not None and costo != producto.precio_costo and not puede_cambiar_costo:
            errores.append((numero_linea, f"Línea {numero_linea} ({codigo}): no tienes permiso para cambiar el costo de los productos."))
            continue

        costo_final = costo if costo is not None else producto.precio_costo
        detalle = []
        for posicion, nuevo in sorted(precios.items()):
            lista, almacen = destinos[posicion]
            antes = precios_actuales.get((producto.pk, lista.pk, almacen.pk if almacen else None))
            if antes == nuevo and costo is None:
                continue
            _, nombre_lista, nombre_almacen = PRECIO_POSICIONES[posicion - 1]
            detalle.append({
                "posicion": posicion,
                "lista": _titulo_posicion(nombre_lista, nombre_almacen),
                "antes": str(antes) if antes is not None else None,
                "nuevo": str(nuevo),
                "utilidad": str(u) if (u := utilidad_de(producto, costo_final, nuevo)) is not None else None,
            })
        cambia_costo = costo is not None and costo != producto.precio_costo
        if not cambia_costo and not detalle:
            continue
        cambios.append({
            "linea": numero_linea,
            "producto_id": producto.pk,
            "sku": producto.sku,
            "nombre": producto.nombre,
            "costo_antes": str(producto.precio_costo),
            "costo_nuevo": str(costo) if cambia_costo else None,
            "precios": detalle,
        })
    errores.sort(key=lambda error: error[0])
    return cambios, [texto for _, texto in errores]


@transaction.atomic
def aplicar(cambios, usuario):
    """Guarda los cambios ya revisados por leer_archivo(). El costo se
    guarda sin el recálculo automático de las listas; cada precio del
    archivo se guarda con su % de utilidad recalculado; al final el precio de
    venta toma el de Público general. Regresa cuántos productos cambiaron."""
    destinos = _destinos_de_posiciones()
    productos = Producto.objects.select_for_update().in_bulk([c["producto_id"] for c in cambios])
    actualizados = 0
    for cambio in cambios:
        producto = productos.get(cambio["producto_id"])
        if producto is None:
            continue
        if cambio["costo_nuevo"] is not None:
            producto.precio_costo = Decimal(cambio["costo_nuevo"])
            producto.recalcular_precios_lista = False
            producto.updated_by = usuario
            producto.save()
        for detalle in cambio["precios"]:
            lista, almacen = destinos[detalle["posicion"]]
            nuevo = Decimal(detalle["nuevo"])
            precio = ProductoPrecio.objects.filter(
                producto=producto, **_filtro_posicion(*PRECIO_POSICIONES[detalle["posicion"] - 1][1:]),
            ).first()
            if precio is None:
                precio = ProductoPrecio(producto=producto, lista_precio=lista, almacen=almacen, created_by=usuario)
            precio.precio_con_impuesto = nuevo
            precio.utilidad_pct = utilidad_de(producto, producto.precio_costo, nuevo)
            precio.updated_by = usuario
            precio.save()
        producto.sincronizar_precio_venta_con_publico()
        actualizados += 1
    return actualizados
