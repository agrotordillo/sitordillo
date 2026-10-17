import datetime
from decimal import Decimal

import openpyxl
from django.contrib.auth.decorators import permission_required
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Font, PatternFill

from apps.core.filtros_producto import aplicar_filtros_producto, leer_filtros_producto
from apps.products import columnas_listado
from apps.products.models import Producto
from apps.products.views.product_views import filtrar_busqueda_producto

RELLENO_GRUPO = PatternFill("solid", fgColor="F3F4F6")


def _valor_excel(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, datetime.datetime):
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    return valor


@permission_required("products.view_producto", raise_exception=True)
def producto_exportar_excel_view(request):
    """Descarga en Excel el mismo listado de Productos que se ve en
    pantalla: respeta la búsqueda y los filtros, y también las columnas
    (cuáles y en qué orden) y la agrupación que el usuario eligió, con una
    fila por grupo que dice cuántos productos tiene."""
    claves_columnas, claves_agrupar = columnas_listado.preferencia_de(request.user)
    columnas = [columnas_listado.COLUMNAS[clave] for clave in claves_columnas]

    queryset = filtrar_busqueda_producto(Producto.objects.all(), request.GET.get("q", "").strip())
    queryset = aplicar_filtros_producto(queryset, leer_filtros_producto(request))
    conteos = columnas_listado.conteos_por_grupo(queryset, claves_agrupar) if claves_agrupar else {}
    queryset = columnas_listado.preparar_queryset(queryset.order_by("nombre"), claves_columnas, claves_agrupar)

    workbook = openpyxl.Workbook()
    hoja = workbook.active
    hoja.title = "Productos"
    hoja.append([columna.titulo for columna in columnas])
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    for fila in columnas_listado.filas_agrupadas(queryset.iterator(), claves_agrupar, conteos):
        if fila["es_grupo"]:
            sangria = "    " * fila["nivel"]
            hoja.append([f'{sangria}{fila["titulo"]}: {fila["valor"]} ({fila["conteo"]})'])
            numero = hoja.max_row
            hoja.cell(row=numero, column=1).font = Font(bold=True)
            for indice in range(1, len(columnas) + 1):
                hoja.cell(row=numero, column=indice).fill = RELLENO_GRUPO
            continue
        hoja.append([_valor_excel(columna.valor(fila["producto"])) for columna in columnas])

    for indice, columna in enumerate(columnas, start=1):
        letra = hoja.cell(row=1, column=indice).column_letter
        hoja.column_dimensions[letra].width = 40 if columna.formato == "nombre" else 18
        if columna.formato == "fecha":
            for (celda,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celda.number_format = "yyyy-mm-dd hh:mm"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    fecha = timezone.localdate().isoformat()
    response["Content-Disposition"] = f'attachment; filename="productos_{fecha}.xlsx"'
    workbook.save(response)
    return response
