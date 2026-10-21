from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse

from apps.products.models import Almacen, ListaPrecio, Producto, ProductoPrecio, PuntoVenta
from apps.products.services import abrir_turno, resolver_precio_linea

User = get_user_model()


class NextSeguroTests(TestCase):
    """B12 (docs/AUDITORIA.md): después de guardar se redirigía a lo que
    trajera `next`, aunque apuntara a otro sitio."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.producto = Producto.objects.create(nombre="Croqueta", sku="CRQ")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_activar_desactivar_no_redirige_fuera_del_sitio(self):
        url = reverse("products:product-toggle-activo", args=[self.producto.pk])

        respuesta = self.client.post(url, {"next": "https://evil.example/robo"})
        self.assertRedirects(respuesta, reverse("products:product-list"), fetch_redirect_response=False)

        respuesta = self.client.post(url, {"next": "/productos/?q=croqueta"})
        self.assertRedirects(respuesta, "/productos/?q=croqueta", fetch_redirect_response=False)

    def test_el_formulario_no_pinta_un_cancelar_hacia_otro_sitio(self):
        respuesta = self.client.get(
            reverse("products:product-update", args=[self.producto.pk]), {"next": "javascript:alert(1)"}
        )
        self.assertEqual(respuesta.context["next_url"], str(reverse("products:product-list")))
        self.assertNotContains(respuesta, "javascript:alert(1)")


class TurnoYaAbiertoTests(TestCase):
    """B27 (docs/AUDITORIA.md): abrir un turno en una caja que ya tenía uno
    mostraba el nombre técnico de la restricción de la base de datos."""

    def test_mensaje_claro(self):
        sucursal = Almacen.objects.create(nombre="Sucursal Centro", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        caja = PuntoVenta.objects.create(almacen=sucursal, codigo="C1", numero=1, nombre="Caja 1")
        admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        otro = User.objects.create_superuser(username="otro", password="S3guridad!2026")
        abrir_turno(caja, otro)

        with self.assertRaises(ValidationError) as error:
            abrir_turno(caja, admin)
        self.assertEqual(error.exception.messages, ["Este punto de venta ya tiene un turno abierto."])

        self.client.force_login(admin)
        respuesta = self.client.post(reverse("products:turno-abrir"), {"punto_venta": caja.pk}, follow=True)
        mensajes = [str(m) for m in respuesta.context["messages"]]
        self.assertIn("Este punto de venta ya tiene un turno abierto.", mensajes)
        self.assertFalse(any("trn_un_turno" in m for m in mensajes))


class PrecioDePaqueteTests(TestCase):
    """B25 (docs/AUDITORIA.md), decisión del usuario: un paquete se arma a
    su costo real y se vende siempre con la lista PROMOCION, sin importar la
    lista del cliente."""

    @classmethod
    def setUpTestData(cls):
        cls.sucursal = Almacen.objects.create(nombre="Sucursal Centro", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        cls.mayoreo = ListaPrecio.objects.get_or_create(nombre="MAYOREO")[0]
        cls.promocion = ListaPrecio.objects.get_or_create(nombre="PROMOCION")[0]
        cls.paquete = Producto.objects.create(
            nombre="Combo croqueta", sku="CMB", tipo=Producto.TipoProducto.PAQUETE, almacen=cls.sucursal,
            precio_venta=Decimal("300.00"),
        )
        cls.croqueta = Producto.objects.create(nombre="Croqueta", sku="CRQ", precio_venta=Decimal("120.00"))
        for producto, lista, precio in (
            (cls.paquete, cls.mayoreo, "280.00"),
            (cls.paquete, cls.promocion, "250.00"),
            (cls.croqueta, cls.mayoreo, "110.00"),
            (cls.croqueta, cls.promocion, "100.00"),
        ):
            ProductoPrecio.objects.create(producto=producto, lista_precio=lista, precio_con_impuesto=Decimal(precio))

    def test_el_paquete_se_cobra_con_promocion_y_lo_demas_con_la_lista_del_cliente(self):
        self.assertEqual(
            resolver_precio_linea(self.paquete, self.mayoreo, self.sucursal), (Decimal("250.00"), self.promocion),
        )
        self.assertEqual(
            resolver_precio_linea(self.croqueta, self.mayoreo, self.sucursal), (Decimal("110.00"), self.mayoreo),
        )

    def test_paquete_sin_precio_de_promocion_usa_la_lista_del_cliente(self):
        ProductoPrecio.objects.filter(producto=self.paquete, lista_precio=self.promocion).delete()
        self.assertEqual(
            resolver_precio_linea(self.paquete, self.mayoreo, self.sucursal), (Decimal("280.00"), self.mayoreo),
        )

    def test_el_buscador_ignora_una_sucursal_de_precio_invalida(self):
        # B28: antes daba error 500.
        self.client.force_login(User.objects.create_superuser(username="admin2", password="S3guridad!2026"))
        respuesta = self.client.get(reverse("api:producto-buscar"), {"q": "Combo", "precio_almacen": "abc"})
        self.assertEqual(respuesta.json()[0]["precio_venta"], "300.00")

    def test_el_buscador_de_ventas_muestra_el_precio_de_promocion(self):
        self.client.force_login(User.objects.create_superuser(username="admin", password="S3guridad!2026"))
        respuesta = self.client.get(reverse("api:producto-buscar"), {"q": "Combo", "precio_almacen": self.sucursal.pk})
        self.assertEqual(respuesta.json()[0]["precio_venta"], "250.00")


class ConfirmarActualizarPreciosListaTests(TestCase):
    """Al cambiar el costo en la edición del producto se pregunta si también
    se actualizan los precios de lista (campo oculto actualizar_precios_lista,
    ver producto-costo-confirmar.js). Si sí, las listas con % de utilidad se
    recalculan y el precio de venta toma el de Público en general."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.publico = ListaPrecio.objects.get_or_create(nombre="PUBLICO", defaults={"orden": 1})[0]
        cls.mayoreo = ListaPrecio.objects.get_or_create(nombre="MAYOREO", defaults={"orden": 3})[0]

    def setUp(self):
        self.client.force_login(self.admin)
        self.producto = Producto.objects.create(
            nombre="Croqueta", sku="CRQ", precio_costo=Decimal("100.0000"), precio_venta=Decimal("151.00"),
            tasa_iva=Decimal("16.00"),
        )
        # Público con 30% de utilidad: 100 × 1.30 × 1.16 = 150.80 -> 151.00.
        self.precio_publico = ProductoPrecio.objects.create(
            producto=self.producto, lista_precio=self.publico, utilidad_pct=Decimal("30"),
            precio_con_impuesto=Decimal("151.00"),
        )
        # Mayoreo capturado a mano (sin % de utilidad): nunca se recalcula.
        self.precio_mayoreo = ProductoPrecio.objects.create(
            producto=self.producto, lista_precio=self.mayoreo, precio_con_impuesto=Decimal("140.00"),
        )

    def _editar(self, **cambios):
        datos = {
            "nombre": self.producto.nombre, "sku": self.producto.sku, "tipo": "producto",
            "tipo_iva": "gravado", "tasa_iva": "16.00", "precio_costo": "100.0000", "precio_venta": "151.00",
            "stock_minimo": "0", "stock_maximo": "0",
        }
        datos.update(cambios)
        return self.client.post(reverse("products:product-update", args=[self.producto.pk]), datos)

    def test_si_actualiza_listas_con_utilidad_y_precio_de_venta_toma_publico(self):
        respuesta = self._editar(precio_costo="110.0000", precio_venta="999.00", actualizar_precios_lista="True")
        self.assertEqual(respuesta.status_code, 302)
        self.precio_publico.refresh_from_db()
        self.precio_mayoreo.refresh_from_db()
        self.producto.refresh_from_db()
        # 110 × 1.30 × 1.16 = 165.88 -> 166.00 (redondeo hacia arriba a $0.50).
        self.assertEqual(self.precio_publico.precio_con_impuesto, Decimal("166.00"))
        self.assertEqual(self.precio_mayoreo.precio_con_impuesto, Decimal("140.00"))
        self.assertEqual(self.producto.precio_costo, Decimal("110.0000"))
        self.assertEqual(self.producto.precio_venta, Decimal("166.00"))

    def test_no_guarda_solo_el_costo(self):
        respuesta = self._editar(precio_costo="110.0000", precio_venta="155.00", actualizar_precios_lista="False")
        self.assertEqual(respuesta.status_code, 302)
        self.precio_publico.refresh_from_db()
        self.producto.refresh_from_db()
        self.assertEqual(self.precio_publico.precio_con_impuesto, Decimal("151.00"))
        self.assertEqual(self.producto.precio_costo, Decimal("110.0000"))
        self.assertEqual(self.producto.precio_venta, Decimal("155.00"))

    def test_sin_respuesta_conserva_el_recalculo_de_siempre(self):
        # Sin JavaScript el campo oculto viaja con el valor con el que se pinta.
        formulario = self.client.get(reverse("products:product-update", args=[self.producto.pk]))
        self.assertContains(formulario, 'name="actualizar_precios_lista" value="True"')
        self._editar(precio_costo="110.0000", actualizar_precios_lista="True")
        self.precio_publico.refresh_from_db()
        self.assertEqual(self.precio_publico.precio_con_impuesto, Decimal("166.00"))

    def test_el_formulario_trae_las_listas_para_la_confirmacion(self):
        respuesta = self.client.get(reverse("products:product-update", args=[self.producto.pk]))
        datos = respuesta.context["precios_con_utilidad"]
        self.assertEqual(datos["publico_general"], "151.00")
        self.assertEqual([f["lista"] for f in datos["filas"]], ["PUBLICO"])
        self.assertTrue(datos["filas"][0]["es_publico_general"])

    def datos_pantalla_precios(self, costo="100.0000", venta="151.00", publico="151.00", mayoreo="140.00"):
        return {
            "producto-precio_costo": costo, "producto-precio_venta": venta,
            "precios-TOTAL_FORMS": "2", "precios-INITIAL_FORMS": "2",
            "precios-MIN_NUM_FORMS": "0", "precios-MAX_NUM_FORMS": "1000",
            "precios-0-id": self.precio_publico.pk, "precios-0-lista_precio": self.publico.pk,
            "precios-0-almacen": "", "precios-0-utilidad_pct": "30", "precios-0-precio_con_impuesto": publico,
            "precios-1-id": self.precio_mayoreo.pk, "precios-1-lista_precio": self.mayoreo.pk,
            "precios-1-almacen": "", "precios-1-utilidad_pct": "", "precios-1-precio_con_impuesto": mayoreo,
        }

    def test_guardar_la_pantalla_de_precios_iguala_el_precio_de_venta_a_publico(self):
        url = reverse("products:producto-precios", args=[self.producto.pk])
        respuesta = self.client.post(url, self.datos_pantalla_precios(venta="160.00", publico="175.00"))
        self.assertEqual(respuesta.status_code, 302)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.precio_venta, Decimal("175.00"))

    def test_pantalla_de_precios_cambia_el_costo_sin_recalcular_solo(self):
        # Sin confirmar: el costo cambia y los precios se guardan tal cual vienen.
        url = reverse("products:producto-precios", args=[self.producto.pk])
        respuesta = self.client.post(url, self.datos_pantalla_precios(costo="120.0000"))
        self.assertEqual(respuesta.status_code, 302)
        self.producto.refresh_from_db()
        self.precio_publico.refresh_from_db()
        self.assertEqual(self.producto.precio_costo, Decimal("120.0000"))
        self.assertEqual(self.precio_publico.precio_con_impuesto, Decimal("151.00"))
        self.assertEqual(self.precio_publico.utilidad_pct, Decimal("30"))

    def test_pantalla_de_precios_guarda_lo_recalculado_al_confirmar(self):
        # Confirmado: la pantalla manda Público ya recalculado
        # (120 × 1.30 × 1.16 = 180.96 -> 181.00) y el precio de venta lo toma.
        url = reverse("products:producto-precios", args=[self.producto.pk])
        respuesta = self.client.post(url, self.datos_pantalla_precios(costo="120.0000", venta="181.00", publico="181.00"))
        self.assertEqual(respuesta.status_code, 302)
        self.producto.refresh_from_db()
        self.assertEqual((self.producto.precio_costo, self.producto.precio_venta), (Decimal("120.0000"), Decimal("181.00")))

    def test_sin_permiso_de_productos_no_cambia_el_costo(self):
        from django.contrib.auth.models import Permission

        capturista = User.objects.create_user(username="precios", password="S3guridad!2026")
        capturista.user_permissions.add(Permission.objects.get(codename="change_productoprecio"))
        self.client.force_login(capturista)
        url = reverse("products:producto-precios", args=[self.producto.pk])
        self.assertContains(self.client.get(url), "necesitas permiso para editar productos")
        respuesta = self.client.post(url, self.datos_pantalla_precios(costo="1.0000"))
        self.assertEqual(respuesta.status_code, 302)
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.precio_costo, Decimal("100.0000"))


class OrdenListaPreciosTests(TestCase):
    """La pantalla de Precios de un producto siempre muestra las listas en el
    orden del sistema anterior: las 5 generales, Público Iquinuapa, Público
    Huimanguillo y luego Público y Mayoreo del CEDIS (identificado por su
    tipo, no por su nombre). Cualquier otra combinación va al final."""

    def test_orden_canonico_aunque_se_hayan_guardado_revueltas(self):
        from apps.products.forms import ProductoPrecioFormSet

        listas = {
            nombre: ListaPrecio.objects.get_or_create(nombre=nombre, defaults={"orden": orden})[0]
            for orden, nombre in enumerate(["PUBLICO", "MEDIO MAYOREO", "MAYOREO", "SUB DISTRIBUIDOR", "PROMOCION"], 1)
        }
        iquinuapa = Almacen.objects.create(nombre="IQUINUAPA", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        huimanguillo = Almacen.objects.create(nombre="HUIMANGUILLO", tipo=Almacen.Tipo.SUCURSAL, numero=2)
        cedis = Almacen.objects.create(nombre="CEDIS BODEGA SUR", tipo=Almacen.Tipo.CEDIS, numero=3)
        bodega_sur = Almacen.objects.create(nombre="BODEGA SUR", tipo=Almacen.Tipo.SUCURSAL, numero=4)
        producto = Producto.objects.create(nombre="Sulfa", sku="3S100")

        revueltas = [
            ("MAYOREO", bodega_sur), ("MAYOREO", cedis), ("PUBLICO", bodega_sur), ("PROMOCION", None),
            ("PUBLICO", cedis), ("PUBLICO", huimanguillo), ("PUBLICO", None), ("SUB DISTRIBUIDOR", None),
            ("PUBLICO", iquinuapa), ("MAYOREO", None), ("MEDIO MAYOREO", None),
        ]
        for lista, almacen in revueltas:
            ProductoPrecio.objects.create(
                producto=producto, lista_precio=listas[lista], almacen=almacen, precio_con_impuesto=Decimal("10"),
            )

        formset = ProductoPrecioFormSet(instance=producto, prefix="precios")
        orden = [
            (f.instance.lista_precio.nombre, f.instance.almacen.nombre if f.instance.almacen_id else None)
            for f in formset.forms if f.instance.pk
        ]
        self.assertEqual(orden, [
            ("PUBLICO", None), ("MEDIO MAYOREO", None), ("MAYOREO", None), ("SUB DISTRIBUIDOR", None),
            ("PROMOCION", None), ("PUBLICO", "IQUINUAPA"), ("PUBLICO", "HUIMANGUILLO"),
            ("PUBLICO", "CEDIS BODEGA SUR"), ("MAYOREO", "CEDIS BODEGA SUR"),
            # Fuera de la tabla del sistema anterior: al final.
            ("PUBLICO", "BODEGA SUR"), ("MAYOREO", "BODEGA SUR"),
        ])


class UnidadAuxiliarTests(TestCase):
    """Unidad auxiliar del sistema anterior: el contenido de cada presentación
    (1 PIEZA = 100 ML), con el que el negocio identifica el producto."""

    @classmethod
    def setUpTestData(cls):
        from apps.products.models import UnidadMedida

        cls.pieza = UnidadMedida.objects.create(nombre="PIEZA", abreviatura="PZA")
        cls.ml = UnidadMedida.objects.create(nombre="MILILITRO", abreviatura="ML")
        cls.kg = UnidadMedida.objects.create(nombre="KILOGRAMO", abreviatura="KG")

    def test_reglas_de_importacion(self):
        from apps.products.management.commands.import_legacy_productos import campos_unidad_auxiliar

        unidades = {1: self.pieza, 2: self.ml}
        fila = {"id_unidad_medida_auxiliar": 2, "factor_unidad_medida_auxiliar": "100.00000000", "control_auxiliar": "NO"}
        campos, sin_unidad = campos_unidad_auxiliar(fila, unidades)
        self.assertEqual(campos, {"unidad_medida_auxiliar": self.ml, "factor_unidad_auxiliar": Decimal("100.0000"), "control_auxiliar": False})
        self.assertFalse(sin_unidad)

        # Factor 0 en el legado = sin capturar: no se trae la unidad suelta.
        fila = {"id_unidad_medida_auxiliar": 1, "factor_unidad_medida_auxiliar": "0.00000000", "control_auxiliar": "SI"}
        campos, sin_unidad = campos_unidad_auxiliar(fila, unidades)
        self.assertEqual(campos, {"unidad_medida_auxiliar": None, "factor_unidad_auxiliar": None, "control_auxiliar": True})
        self.assertFalse(sin_unidad)

        # Factor sin unidad: no se trae y se reporta.
        fila = {"id_unidad_medida_auxiliar": None, "factor_unidad_medida_auxiliar": "40.00000000", "control_auxiliar": ""}
        campos, sin_unidad = campos_unidad_auxiliar(fila, unidades)
        self.assertIsNone(campos["factor_unidad_auxiliar"])
        self.assertTrue(sin_unidad)

    def test_contenido_display(self):
        producto = Producto(nombre="Sulfa", sku="3S100", unidad_medida_auxiliar=self.ml, factor_unidad_auxiliar=Decimal("100.0000"))
        self.assertEqual(producto.contenido_display, "100 ML")
        producto.unidad_medida_auxiliar, producto.factor_unidad_auxiliar = self.kg, Decimal("0.5000")
        self.assertEqual(producto.contenido_display, "0.5 KG")
        self.assertEqual(Producto(nombre="Sin", sku="X").contenido_display, "")

    def test_formulario_pide_unidad_y_factor_juntos(self):
        admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        self.client.force_login(admin)
        producto = Producto.objects.create(nombre="Sulfa", sku="3S100")
        url = reverse("products:product-update", args=[producto.pk])
        base = {
            "nombre": "Sulfa", "sku": "3S100", "tipo": "producto", "tipo_iva": "gravado", "tasa_iva": "16.00",
            "precio_costo": "0", "precio_venta": "0", "stock_minimo": "0", "stock_maximo": "0",
            "actualizar_precios_lista": "False",
        }

        respuesta = self.client.post(url, {**base, "factor_unidad_auxiliar": "100"})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("unidad_medida_auxiliar", respuesta.context["form"].errors)

        respuesta = self.client.post(url, {**base, "unidad_medida_auxiliar": self.ml.pk})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("factor_unidad_auxiliar", respuesta.context["form"].errors)

        respuesta = self.client.post(url, {**base, "unidad_medida_auxiliar": self.ml.pk, "factor_unidad_auxiliar": "100"})
        self.assertEqual(respuesta.status_code, 302)
        producto.refresh_from_db()
        self.assertEqual(producto.contenido_display, "100 ML")


class ListadoPersonalizableTests(TestCase):
    """"Columnas" y "Agrupar por" del listado de Productos, como en el
    sistema anterior: cada usuario elige qué columnas ver, en qué orden, y
    si se agrupa (estático: un campo; dinámico: varios niveles). Se guarda
    por usuario y lo respeta también el Excel."""

    @classmethod
    def setUpTestData(cls):
        from apps.products.models import Linea, Marca

        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.otro = User.objects.create_superuser(username="otro", password="S3guridad!2026")
        cls.perros = Linea.objects.create(nombre="PERROS")
        cls.gatos = Linea.objects.create(nombre="GATOS")
        cls.marca = Marca.objects.create(nombre="NUPEC")
        publico = ListaPrecio.objects.get_or_create(nombre="PUBLICO", defaults={"orden": 1})[0]
        iquinuapa = Almacen.objects.create(nombre="IQUINUAPA", tipo=Almacen.Tipo.SUCURSAL, numero=1)
        # 30 de PERROS (más de una página), 1 de GATOS y 1 sin línea.
        for numero in range(30):
            Producto.objects.create(nombre=f"Croqueta {numero:02d}", sku=f"CRQ{numero:02d}", linea=cls.perros, marca=cls.marca)
        cls.arena = Producto.objects.create(nombre="Arena", sku="ARN", linea=cls.gatos)
        Producto.objects.create(nombre="Collar", sku="COL")
        ProductoPrecio.objects.create(
            producto=cls.arena, lista_precio=publico, almacen=iquinuapa, precio_con_impuesto=Decimal("88.50"),
        )

    def setUp(self):
        self.client.force_login(self.admin)
        self.url = reverse("products:product-list")
        self.url_preferencias = reverse("products:product-list-preferencias")

    def titulos(self, respuesta):
        return [columna.titulo for columna in respuesta.context["columnas"]]

    def test_columnas_por_defecto(self):
        respuesta = self.client.get(self.url)
        self.assertEqual(
            self.titulos(respuesta),
            ["Código", "Nombre", "Tipo", "Unidad", "Estado", "Costo",
             "Público", "Medio mayoreo", "Mayoreo", "Sub distribuidor", "Promoción"],
        )

    def test_elegir_columnas_y_orden_por_usuario(self):
        respuesta = self.client.post(self.url_preferencias, {
            "accion": "columnas", "columnas": ["nombre", "precio_6", "linea", "no_existe", "nombre"],
            "next": "/productos/?q=arena",
        })
        self.assertRedirects(respuesta, "/productos/?q=arena", fetch_redirect_response=False)

        respuesta = self.client.get(self.url, {"q": "arena"})
        # Costo es fija: si no se manda, se agrega igual.
        self.assertEqual(self.titulos(respuesta), ["Nombre", "Público Iquinuapa", "Línea", "Costo"])
        self.assertContains(respuesta, "88.50")
        self.assertContains(respuesta, "GATOS")

        # El otro usuario sigue con las columnas por defecto.
        self.client.force_login(self.otro)
        self.assertEqual(self.titulos(self.client.get(self.url))[0], "Código")

        self.client.force_login(self.admin)
        self.client.post(self.url_preferencias, {"accion": "restablecer_columnas"})
        self.assertEqual(self.titulos(self.client.get(self.url))[0], "Código")

    def test_agrupar_estatico_cuenta_todo_el_resultado(self):
        self.client.post(self.url_preferencias, {"accion": "agrupar", "modo": "estatico", "agrupar_estatico": "linea"})
        respuesta = self.client.get(self.url)
        grupos = [fila for fila in respuesta.context["filas"] if fila["es_grupo"]]
        # Página 1 (25 productos): GATOS primero, luego PERROS con sus 30 aunque
        # en esta página solo se vean 24.
        self.assertEqual([(g["valor"], g["conteo"]) for g in grupos], [("GATOS", 1), ("PERROS", 30)])

        respuesta = self.client.get(self.url, {"page": 2})
        grupos = [(f["valor"], f["conteo"]) for f in respuesta.context["filas"] if f["es_grupo"]]
        self.assertEqual(grupos, [("PERROS", 30), ("Sin línea", 1)])

        # Un valor que no está en la lista de estáticos no agrupa.
        self.client.post(self.url_preferencias, {"accion": "agrupar", "modo": "estatico", "agrupar_estatico": "tipo"})
        self.assertEqual(self.client.get(self.url).context["agrupar"], [])

    def test_agrupar_dinamico_anida_niveles(self):
        self.client.post(self.url_preferencias, {
            "accion": "agrupar", "modo": "dinamico", "agrupar": ["linea", "marca", "inventado"],
        })
        respuesta = self.client.get(self.url, {"q": "croqueta"})
        filas = respuesta.context["filas"]
        self.assertEqual([a.clave for a in respuesta.context["agrupar"]], ["linea", "marca"])
        linea, marca = filas[0], filas[1]
        self.assertEqual((linea["valor"], linea["nivel"], linea["conteo"]), ("PERROS", 0, 30))
        self.assertEqual((marca["valor"], marca["nivel"], marca["grupos"]), ("NUPEC", 1, linea["id"]))
        self.assertEqual(filas[2]["grupos"], f'{linea["id"]} {marca["id"]}')

    def test_excel_respeta_columnas_y_grupos(self):
        import io

        import openpyxl

        self.client.post(self.url_preferencias, {"accion": "columnas", "columnas": ["sku", "nombre"]})
        self.client.post(self.url_preferencias, {"accion": "agrupar", "modo": "estatico", "agrupar_estatico": "linea"})
        respuesta = self.client.get(reverse("products:product-exportar"))
        hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
        filas = [tuple(celda.value for celda in fila) for fila in hoja.iter_rows()]
        self.assertEqual(filas[0], ("Código", "Nombre", "Costo"))
        self.assertEqual(filas[1][0], "Línea: GATOS (1)")
        self.assertEqual(filas[2], ("ARN", "Arena", 0.0))
        self.assertEqual(filas[3][0], "Línea: PERROS (30)")


class BusquedaProductoTests(TestCase):
    """Todos los buscadores de producto buscan por nombre y SKU (código):
    palabra por palabra, con el SKU exacto primero, y con 1-2 caracteres
    solo el SKU exacto (hay SKUs de una o dos letras)."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.sulfa = Producto.objects.create(nombre="3 SULFA INYECTABLE 100ML", sku="3S100")
        cls.cm = Producto.objects.create(nombre="CADENA METALICA", sku="CM")
        cls.cm_largo = Producto.objects.create(nombre="ACM CEPILLO", sku="ACM1")
        cls.paquete = Producto.objects.create(nombre="COMBO SULFA", sku="CMB1", tipo=Producto.TipoProducto.PAQUETE)

    def setUp(self):
        self.client.force_login(self.admin)

    def buscar(self, q, **extra):
        respuesta = self.client.get(reverse("api:producto-buscar"), {"q": q, **extra})
        return [item["sku"] for item in respuesta.json()]

    def test_por_palabras_en_cualquier_orden(self):
        self.assertEqual(self.buscar("sulfa 100"), ["3S100"])
        self.assertEqual(self.buscar("100ml inyectable"), ["3S100"])
        self.assertEqual(self.buscar("sulfa 250"), [])

    def test_sku_corto_solo_exacto_y_exacto_primero(self):
        self.assertEqual(self.buscar("cm"), ["CM"])
        self.assertEqual(self.buscar("c"), [])
        # Con 3+ caracteres también por nombre/SKU parcial, el exacto primero.
        self.assertEqual(self.buscar("acm1")[0], "ACM1")
        self.assertEqual(self.buscar("3s100"), ["3S100"])

    def test_filtro_por_tipo(self):
        self.assertEqual(self.buscar("sulfa"), ["3S100", "CMB1"])
        self.assertEqual(self.buscar("sulfa", tipo="paquete"), ["CMB1"])
        self.assertEqual(self.buscar("sulfa", tipo="sin_paquete"), ["3S100"])

    def test_listado_usa_la_misma_busqueda(self):
        respuesta = self.client.get(reverse("products:product-list"), {"q": "inyectable sulfa"})
        self.assertEqual([p.sku for p in respuesta.context["products"]], ["3S100"])

    def test_formularios_sin_select_de_productos(self):
        paginas = [
            reverse("inventario:receta-conversion-create"),
            reverse("inventario:ensamble-create"),
            reverse("products:paquete-componentes", args=[self.paquete.pk]),
        ]
        for url in paginas:
            with self.subTest(url=url):
                respuesta = self.client.get(url)
                self.assertEqual(respuesta.status_code, 200)
                self.assertContains(respuesta, "producto-search-input")
                self.assertNotContains(respuesta, "CADENA METALICA</option>")


class PreciosMasivosTests(TestCase):
    """"Actualizar precios" con la plantilla del sistema anterior: se revisa
    el archivo (qué cambia y errores por línea) sin guardar nada, y solo se
    aplica al confirmar. Celdas vacías o en 0 no cambian; el % de utilidad
    de cada precio capturado se recalcula; los no capturados no se tocan."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.publico = ListaPrecio.objects.get_or_create(nombre="PUBLICO", defaults={"orden": 1})[0]
        cls.mayoreo = ListaPrecio.objects.get_or_create(nombre="MAYOREO", defaults={"orden": 3})[0]
        cls.url = reverse("products:precios-masivos")

    def setUp(self):
        self.client.force_login(self.admin)
        # Exento de IVA para que las cuentas sean directas.
        self.sulfa = Producto.objects.create(
            nombre="3 SULFA", sku="3S100", precio_costo=Decimal("80.0000"), precio_venta=Decimal("115.00"),
            tipo_iva=Producto.TipoIVA.EXENTO,
        )
        self.publico_sulfa = ProductoPrecio.objects.create(
            producto=self.sulfa, lista_precio=self.publico, utilidad_pct=Decimal("43.75"), precio_con_impuesto=Decimal("115.00"),
        )
        self.mayoreo_sulfa = ProductoPrecio.objects.create(
            producto=self.sulfa, lista_precio=self.mayoreo, utilidad_pct=Decimal("28.125"), precio_con_impuesto=Decimal("102.50"),
        )
        self.eje = Producto.objects.create(nombre="EJE SOLO", sku="E", precio_costo=Decimal("10.0000"), tipo_iva=Producto.TipoIVA.EXENTO)

    def archivo(self, renglones, titulo=None):
        import io

        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.products.precios_masivos import generar_plantilla

        libro = generar_plantilla()
        hoja = libro["Precios"]
        if titulo is not None:
            hoja["A1"] = titulo
        for renglon in renglones:
            hoja.append(renglon)
        contenido = io.BytesIO()
        libro.save(contenido)
        return SimpleUploadedFile("precios.xlsx", contenido.getvalue())

    def revisar(self, renglones, **kwargs):
        return self.client.post(self.url, {"accion": "revisar", "archivo": self.archivo(renglones, **kwargs)})

    def test_plantilla_como_la_del_sistema_anterior(self):
        import io

        import openpyxl

        respuesta = self.client.get(reverse("products:precios-masivos-plantilla"))
        libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
        hoja = libro["Precios"]
        self.assertEqual(hoja["A1"].value, "Actualización de precios")
        self.assertEqual(
            [c.value for c in hoja[2]],
            ["codigo_producto", "nombre_producto", "costo"] + [f"precio_{n}" for n in range(1, 10)],
        )
        self.assertIn("Instrucciones", libro.sheetnames)

    def test_revisar_no_guarda_y_aplicar_si(self):
        # Costo 90 y Público 130 (Mayoreo vacío: no se toca); "e" en minúsculas.
        respuesta = self.revisar([
            ["3S100", "3 SULFA", 90, 130, None, 0],
            ["e", "EJE", None, None, None, 25],
        ])
        cambios = respuesta.context["cambios"]
        self.assertEqual(respuesta.context["errores"], [])
        self.assertEqual([c["sku"] for c in cambios], ["3S100", "E"])
        self.assertEqual(cambios[0]["costo_nuevo"], "90.0000")
        self.assertEqual([(p["lista"], p["antes"], p["nuevo"], p["utilidad"]) for p in cambios[0]["precios"]],
                         [("Público", "115.00", "130.00", "44.4444")])
        self.assertEqual([(p["lista"], p["antes"], p["nuevo"]) for p in cambios[1]["precios"]], [("Mayoreo", None, "25.00")])
        self.sulfa.refresh_from_db()
        self.assertEqual(self.sulfa.precio_costo, Decimal("80.0000"))

        respuesta = self.client.post(self.url, {"accion": "aplicar", "token": respuesta.context["token"]})
        self.assertRedirects(respuesta, self.url, fetch_redirect_response=False)
        self.sulfa.refresh_from_db()
        self.publico_sulfa.refresh_from_db()
        self.mayoreo_sulfa.refresh_from_db()
        self.assertEqual((self.sulfa.precio_costo, self.sulfa.precio_venta), (Decimal("90.0000"), Decimal("130.00")))
        self.assertEqual((self.publico_sulfa.precio_con_impuesto, self.publico_sulfa.utilidad_pct), (Decimal("130.00"), Decimal("44.4444")))
        self.assertEqual((self.mayoreo_sulfa.precio_con_impuesto, self.mayoreo_sulfa.utilidad_pct), (Decimal("102.50"), Decimal("28.1250")))
        nuevo = ProductoPrecio.objects.get(producto=self.eje, lista_precio=self.mayoreo, almacen__isnull=True)
        self.assertEqual((nuevo.precio_con_impuesto, nuevo.utilidad_pct), (Decimal("25.00"), Decimal("150.0000")))

        # La misma revisión no se aplica dos veces.
        respuesta = self.client.post(self.url, {"accion": "aplicar", "token": "otro"})
        self.assertRedirects(respuesta, self.url, fetch_redirect_response=False)

    def test_errores_por_linea_y_no_se_aplica_nada(self):
        respuesta = self.revisar([
            ["NOEXISTE", "", 10],
            ["3S100", "", -5],
            ["E", "", None, "abc"],
            ["E", "", 12],
            ["E", "", 13],
        ])
        errores = respuesta.context["errores"]
        self.assertEqual(len(errores), 4)
        self.assertIn("Línea 3: el código de producto NOEXISTE no existe.", errores)
        self.assertTrue(errores[1].startswith("Línea 4 (3S100): el costo no puede ser negativo"))
        self.assertTrue(errores[2].startswith("Línea 5 (E): precio_1"))
        self.assertIn("Línea 7: el código E ya viene en la línea 6.", errores)
        self.assertNotIn("precios_masivos_revision", self.client.session)
        self.assertNotContains(respuesta, "Aplicar cambios")

        respuesta = self.revisar([["3S100", "", 90]], titulo="Otra cosa")
        self.assertIn("estructura de la plantilla", respuesta.context["errores"][0])

    def test_sin_permiso_de_productos_no_cambia_el_costo(self):
        from django.contrib.auth.models import Permission

        capturista = User.objects.create_user(username="precios", password="S3guridad!2026")
        capturista.user_permissions.add(Permission.objects.get(codename="change_productoprecio"))
        self.client.force_login(capturista)
        respuesta = self.revisar([["3S100", "", 90, 130], ["E", "", None, 20]])
        self.assertEqual(len(respuesta.context["errores"]), 1)
        self.assertIn("no tienes permiso para cambiar el costo", respuesta.context["errores"][0])


class SubcategoriaIndependienteTests(TestCase):
    """La subcategoría es un catálogo independiente (como Marca o Clase): ya
    no cuelga de una categoría, y el producto tiene las dos por separado."""

    @classmethod
    def setUpTestData(cls):
        from apps.products.models import Categoria, Subcategoria

        cls.admin = User.objects.create_superuser(username="admin", password="S3guridad!2026")
        cls.perros = Categoria.objects.create(nombre="PERROS")
        cls.gatos = Categoria.objects.create(nombre="GATOS")
        cls.alimento = Subcategoria.objects.create(nombre="ALIMENTO SECO")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_cualquier_combinacion_de_categoria_y_subcategoria(self):
        for categoria in (self.perros, self.gatos):
            producto = Producto(nombre=f"Croqueta {categoria.nombre}", sku=f"CRQ-{categoria.nombre}", categoria=categoria, subcategoria=self.alimento)
            producto.full_clean()
            producto.save()
        respuesta = self.client.get(reverse("products:product-list"), {"subcategoria": self.alimento.pk})
        self.assertEqual(len(respuesta.context["products"]), 2)
        # El filtro ofrece todas las subcategorías aunque no se elija categoría.
        self.assertContains(respuesta, f'<option value="{self.alimento.pk}" selected>ALIMENTO SECO</option>', html=True)

    def test_alta_y_nombre_unico(self):
        respuesta = self.client.post(reverse("products:subcategory-create"), {"nombre": "MEDICAMENTO", "descripcion": ""})
        self.assertRedirects(respuesta, reverse("products:subcategory-list"), fetch_redirect_response=False)
        respuesta = self.client.post(reverse("products:subcategory-create"), {"nombre": "MEDICAMENTO", "descripcion": ""})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("nombre", respuesta.context["form"].errors)

    def test_alta_rapida_desde_el_formulario_de_producto(self):
        respuesta = self.client.post(reverse("api:subcategoria-quick-create"), {"nombre": "ACCESORIOS"})
        self.assertEqual(respuesta.status_code, 201)
        self.assertEqual(respuesta.json()["label"], "ACCESORIOS")
        self.assertContains(self.client.get(reverse("products:product-create")), "ACCESORIOS")
